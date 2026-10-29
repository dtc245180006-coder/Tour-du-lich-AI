"""Complete the retained user-guide template with verified VietJourney workflows."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from docx.shared import Emu
import hashlib
import json

ROOT = Path(__file__).resolve().parents[1]
REF = Path(r'C:\Users\OMEN\Downloads\07_GenAI_SoftwareDevelopment_user-guide.docx')
OUT = ROOT/'docs/07_GenAI_SoftwareDevelopment_user-guide.docx'
QA = ROOT/'artifacts/user-guide'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'w':W,'r':R}
def tag(name): return '{'+W+'}'+name
def ser(x): return E.tostring(x, encoding='UTF-8', xml_declaration=True, standalone=True)
with ZipFile(REF) as z:
    parts = {n:z.read(n) for n in z.namelist()}
root = E.fromstring(parts['word/document.xml'])
body = root.find('w:body', NS)
source_nodes = list(body)
heading1 = deepcopy(source_nodes[6])
heading2 = deepcopy(source_nodes[9])
normal = deepcopy(source_nodes[10])
rels = E.fromstring(parts['word/_rels/document.xml.rels'])
toc = source_nodes[4].find('w:sdtContent', NS)
toc_title = deepcopy(toc[0])
toc1 = deepcopy(toc.xpath('w:p[w:pPr/w:pStyle[@w:val="TOC1"]]',namespaces=NS)[0])
toc2 = deepcopy(toc.xpath('w:p[w:pPr/w:pStyle[@w:val="TOC2"]]',namespaces=NS)[0])
section = deepcopy(source_nodes[-1])
for x in source_nodes[6:]: body.remove(x)
entries=[]
picture_count=0

def clean_p(template, text, bold=False):
    p = deepcopy(template)
    for node in list(p):
        if node.tag != tag('pPr'): p.remove(node)
    pp = p.find('w:pPr', NS)
    if pp is None: pp = E.SubElement(p, tag('pPr'))
    # Blank slots are blue italic drafting hints; finished prose uses the same TNR 13 font.
    if template is normal:
        for name in ['ind', 'rPr']:
            node=pp.find('w:'+name,NS)
            if node is not None: pp.remove(node)
    run=E.SubElement(p,tag('r'))
    rp=E.SubElement(run,tag('rPr'))
    E.SubElement(rp,tag('rFonts'), {tag('ascii'):'Times New Roman',tag('hAnsi'):'Times New Roman',tag('cs'):'Times New Roman'})
    E.SubElement(rp,tag('sz'), {tag('val'):'26'})
    E.SubElement(rp,tag('szCs'), {tag('val'):'26'})
    E.SubElement(rp,tag('color'), {tag('val'):'000000'})
    if bold: E.SubElement(rp,tag('b'))
    E.SubElement(run,tag('t')).text=text
    return p

def p(text,bold=False):
    node=clean_p(normal,text,bold)
    body.append(node)
    return node

def h(text,level=1):
    node=clean_p(heading1 if level==1 else heading2,text,True)
    pp=node.find('w:pPr',NS)
    if pp.find('w:keepNext',NS) is None:E.SubElement(pp,tag('keepNext'))
    if level==1 and entries:E.SubElement(pp,tag('pageBreakBefore'))
    index=len(entries)+100
    anchor='UserGuide'+str(index)
    node.insert(1,E.Element(tag('bookmarkStart'),{tag('id'):str(index),tag('name'):anchor}))
    node.append(E.Element(tag('bookmarkEnd'),{tag('id'):str(index)}))
    body.append(node)
    entries.append((text,level,anchor))
    return node

def sub(text):
    node=p(text,True)
    E.SubElement(node.find('w:pPr',NS),tag('keepNext'))

def steps(*items):
    for i,text in enumerate(items,1):p(f'Bước {i}. {text}')

def figure(filename,caption):
    global picture_count
    picture_count+=1
    path=QA/'screenshots'/filename
    temp=Document()
    temp.add_picture(str(path),width=Emu(5_760_000))
    drawing=E.fromstring(E.tostring(temp.element.xpath('//w:drawing')[0]))
    rid=f'rIdGuideImage{picture_count}'
    drawing.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip').set('{'+R+'}embed',rid)
    prop=drawing.find('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')
    prop.set('id',str(1000+picture_count)); prop.set('name',caption); prop.set('descr',caption)
    node=clean_p(normal,'')
    pp=node.find('w:pPr',NS)
    E.SubElement(pp,tag('jc'),{tag('val'):'center'})
    E.SubElement(pp,tag('keepNext'))
    node.find('w:r',NS).append(drawing)
    body.append(node)
    cap=p(f'Hình {picture_count}. {caption}')
    E.SubElement(cap.find('w:pPr',NS),tag('jc'),{tag('val'):'center'})
    parts[f'word/media/user_guide_{picture_count:02d}.png']=path.read_bytes()
    E.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=rid,Type=R+'/image',Target=f'media/user_guide_{picture_count:02d}.png')

h('GIỚI THIỆU ỨNG DỤNG')
p('VietJourney AI là hệ thống quản lý tour du lịch có tích hợp AI, hỗ trợ giới thiệu tour, tiếp nhận yêu cầu tư vấn, quản lý lịch khởi hành, đặt chỗ, thanh toán, hướng dẫn viên và phản hồi khách hàng. Tài liệu hướng dẫn thao tác trên phiên bản mã nguồn hiện tại, dành cho khách vãng lai và nhân viên vận hành.')
p('Quy trình tổng quát: khách chọn tour và gửi yêu cầu → tư vấn viên liên hệ, duyệt và giữ chỗ → kế toán ghi nhận tiền → quản trị viên điều hành chuyến → theo dõi phản hồi và báo cáo. Khách gửi yêu cầu chưa được xem là đã giữ chỗ hoặc đã thanh toán.')
p('Khách vãng lai không cần tài khoản để xem tour, xem hướng dẫn viên, sử dụng tư vấn AI và gửi yêu cầu. Ba vai trò đăng nhập gồm Quản trị viên (ADMIN), Nhân viên tư vấn (CONSULTANT) và Kế toán (ACCOUNTANT). Hướng dẫn viên được quản lý dưới dạng hồ sơ; phiên bản hiện tại chưa có cổng đăng nhập riêng cho hướng dẫn viên.')
p('Địa chỉ chạy mặc định trên máy chủ: http://127.0.0.1:8000. Ảnh minh họa trong tài liệu được chụp từ môi trường demo độc lập; tên khách, số liệu, ngày khởi hành và trạng thái có thể khác dữ liệu thực tế.')
figure('01-home.png','Trang chủ VietJourney AI trong môi trường demo')

h('CẤU HÌNH PHẦN CỨNG - PHẦN MỀM')
h('Phần cứng',2)
p('Cấu hình đề xuất cho việc học tập và trình diễn cục bộ: máy tính có bộ xử lý từ 2 nhân, RAM từ 4 GB, dung lượng trống khoảng 2 GB cho mã nguồn, môi trường Python và dữ liệu demo; màn hình từ 1366 × 768 để thao tác bảng thuận tiện. Đây là mức đề xuất, chưa phải kết quả đo tải hay cấu hình bảo đảm cho nhiều người dùng đồng thời.')
p('Máy chủ cục bộ và trình duyệt có thể chạy trên cùng máy. Kết nối Internet cần thiết khi cài thư viện hoặc sử dụng Gemini; những nghiệp vụ thông thường và chức năng AI nội bộ có thể hoạt động khi không cấu hình khóa AI cloud.')
h('Phần mềm',2)
p('Môi trường sử dụng: Windows, Python 3.10 trở lên và trình duyệt hỗ trợ JavaScript như Chrome hoặc Edge. Hệ thống dùng FastAPI và SQLAlchemy; mặc định có thể sử dụng SQLite. MySQL là lựa chọn triển khai khi đã có máy chủ cơ sở dữ liệu và cấu hình kết nối phù hợp. Không cần cài Node.js để sử dụng bản giao diện đã đóng gói.')
sub('Cài đặt và khởi chạy lần đầu')
steps('Mở thư mục dự án, mở PowerShell tại thư mục đó. Tạo môi trường Python bằng lệnh: py -m venv .venv.',
      'Cài thư viện bằng lệnh: .venv\\Scripts\\python.exe -m pip install -r requirements-lock.txt.',
      'Nếu chưa có .env, sao chép từ .env.example rồi điều chỉnh cấu hình cần thiết. Giữ nguyên .env đang sử dụng nếu dự án đã được cấu hình; không đưa mật khẩu hoặc khóa API vào tài liệu chia sẻ.',
      'Với cơ sở dữ liệu mới dùng để học tập, chạy: .venv\\Scripts\\python.exe scripts\\setup_database.py --demo. Chỉ dùng lựa chọn này cho dữ liệu demo. Nếu triển khai thực tế, tạo tài khoản riêng bằng: .venv\\Scripts\\python.exe scripts\\setup_database.py --admin ten_quan_tri, rồi nhập mật khẩu theo yêu cầu.',
      'Chạy start.bat. Khi cửa sổ máy chủ báo Application startup complete, mở http://127.0.0.1:8000. Giữ cửa sổ máy chủ hoạt động trong suốt thời gian sử dụng.',
      'Để dừng ứng dụng, quay lại cửa sổ máy chủ và nhấn Ctrl+C. Khi cần chạy lại, mở start.bat rồi chờ thông báo sẵn sàng.')
p('Nếu đã có cơ sở dữ liệu và tài khoản hợp lệ, chỉ cần thực hiện bước khởi chạy. Không chạy lại quy trình tạo dữ liệu mẫu để thay thế dữ liệu đang vận hành. Địa chỉ 127.0.0.1 chỉ dùng trên chính máy chạy ứng dụng; khi triển khai cho máy khác, quản trị viên cần cung cấp địa chỉ truy cập phù hợp.')
sub('Tài khoản demo')
p('Sau khi chủ động khởi tạo dữ liệu demo: quản trị viên dùng admin / admin123; tư vấn viên dùng tuvan / tuvan123; kế toán dùng ketoan / ketoan123. Các tài khoản này chỉ phục vụ trình diễn. Với dữ liệu thực tế, dùng tài khoản riêng do quản trị viên cấp.')
sub('Cấu hình tùy chọn cho quản trị viên')
p('Để dùng MySQL, cấu hình DB_TYPE, DB_HOST, DB_PORT, DB_USER, DB_PASSWORD và DB_NAME hoặc DATABASE_URL trong .env; tạo database trước khi chạy ứng dụng. Sao lưu trước khi chuyển đổi hoặc sửa cấu trúc dữ liệu. Người dùng nghiệp vụ không cần thay đổi phần này.')
p('AI cloud được cấu hình trong giao diện bằng tài khoản Admin hoặc qua biến môi trường. Khi triển khai thực tế, đặt APP_ENV=production và SECRET_KEY ngẫu nhiên dài ít nhất 32 ký tự theo README; chỉ công bố dịch vụ sau khi đã cấu hình bảo mật và sao lưu. Tài liệu này tập trung vào thao tác cục bộ.')
p('Để gửi thư xác nhận, thanh toán, hủy đơn và khôi phục mã đặt tour, người phụ trách kỹ thuật cần cấu hình SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD, SMTP_FROM_NAME, SMTP_FROM_EMAIL và APP_BASE_URL. APP_BASE_URL phải là địa chỉ khách có thể truy cập. Nếu chưa cấu hình SMTP hoặc gửi thư thất bại, không coi thông báo trên màn hình là bằng chứng khách đã nhận được email.')

h('CÁC CHỨC NĂNG CHÍNH')
p('Các bước dưới đây sử dụng tên chức năng trên thanh điều hướng. Nút thêm, sửa, xóa, thu tiền hoặc quản lý tài khoản chỉ khả dụng theo quyền được cấp. Khi không thấy một nút, kiểm tra vai trò hiện tại trước khi thao tác tiếp.')
sub('Đăng nhập và đăng xuất dùng chung cho nhân viên')
steps('Chọn Đăng nhập trên giao diện hoặc mở /login. Nhập tên đăng nhập và mật khẩu do quản trị viên cấp.',
      'Nhấn nút đăng nhập. Khi thành công, kiểm tra họ tên và vai trò hiển thị trước khi vào màn hình nghiệp vụ.',
      'Nếu sai thông tin, nhập lại; nếu tài khoản bị khóa hoặc quên mật khẩu, liên hệ quản trị viên để kiểm tra/cấp lại. Phiên bản hiện tại chưa có luồng tự khôi phục mật khẩu bằng email.',
      'Kết thúc phiên làm việc bằng nút Đăng xuất ở khu vực tài khoản hoặc cuối thanh điều hướng. Nếu phiên hết hạn, đăng nhập lại rồi kiểm tra dữ liệu trước khi gửi lại thao tác đang làm.')
figure('03-login.png','Màn hình đăng nhập và các vai trò demo')

h('Chức năng của Khách vãng lai',2)
sub('Tra cứu tour và lịch khởi hành')
steps('Mở Trang Chủ Tour. Duyệt danh sách tour và danh mục hiển thị; mở chi tiết tour để xem điểm đến, thời lượng, mô tả và lịch trình.',
      'Kiểm tra lịch khởi hành còn mở, ngày đi, ngày về, giá và số chỗ còn. Giá áp dụng khi đặt căn cứ vào lịch khởi hành được chọn.',
      'Nếu chưa biết chọn tour nào, mở Trợ lý AI Gemini và nhập nhu cầu cụ thể, ví dụ: “Tôi muốn đi Nha Trang 3 ngày, ngân sách khoảng 5 triệu/người”. Đối chiếu gợi ý với lịch và giá đang có trước khi gửi yêu cầu.')
sub('Gửi yêu cầu tư vấn và đăng ký tour')
steps('Chọn Gửi Yêu Cầu Tư Vấn Tour hoặc nút yêu cầu tương ứng trên tour. Chọn đúng đợt khởi hành.',
      'Nhập họ tên, số điện thoại, email nếu có, số lượng khách là số nguyên dương và yêu cầu đặc biệt. Kiểm tra lại ngày đi và số người.',
      'Gửi yêu cầu và chờ thông báo xác nhận. Lưu mã yêu cầu hệ thống trả về để đối chiếu với nhân viên.',
      'Chờ tư vấn viên liên hệ. Yêu cầu ở trạng thái chờ xác nhận chưa trừ chỗ và chưa phát sinh thu tiền; chỉ khi nhân viên duyệt thì hệ thống mới giữ chỗ nếu còn đủ số chỗ.')
figure('02-request.png','Biểu mẫu gửi yêu cầu tư vấn tour của khách')
p('Nếu không có lịch để chọn, hiện chưa có đợt mở bán đáp ứng điều kiện. Nếu mất kết nối sau khi gửi, kiểm tra thông báo và mã đã nhận trước khi gửi lại để tránh tạo nhiều yêu cầu. Khách không cần tài khoản để tra cứu đơn; khi muốn hủy hoặc điều chỉnh, liên hệ nhân viên để xử lý.')
sub('Tra cứu đơn tour và hướng dẫn chuyển khoản')
steps('Trên Trang Chủ Tour, chọn Tra Cứu Đơn Tour. Nhập mã đặt chỗ và số điện thoại hoặc email trùng với thông tin đã đăng ký, rồi nhấn Tra Cứu Ngay.',
      'Kiểm tra tour, ngày đi, số người, trạng thái đơn, tổng tiền, tiền đã đóng, công nợ và lịch sử giao dịch hiển thị.',
      'Nếu còn khoản phải thanh toán, xem phần hướng dẫn chuyển khoản/VietQR. Xác minh thông tin người thụ hưởng với đơn vị vận hành trước khi chuyển; dùng đúng nội dung gắn với mã đơn. Đối với yêu cầu còn chờ xác nhận, liên hệ tư vấn viên trước khi chuyển cọc.',
      'Sau khi chuyển, chờ Kế toán xác minh và ghi nhận rồi tra cứu lại. Quét QR hoặc mở ứng dụng ngân hàng không tự cập nhật trạng thái đã thanh toán trong hệ thống.',
      'Nếu quên mã, chọn Quên mã đặt tour? Gửi lại qua Gmail; nhập email đã đăng ký và yêu cầu gửi lại. Kiểm tra hộp thư và mục Spam. Có thể mở liên kết tra cứu trực tiếp trong email khi chức năng gửi thư đã được cấu hình.')
figure('14-tracking.png','Tra cứu đơn tour bằng mã đặt chỗ và thông tin liên hệ')
p('Liên kết tra cứu trong email chứa mã truy cập riêng của đơn. Không chuyển tiếp liên kết cho người không liên quan. Nếu không nhận được thư, kiểm tra lại email đăng ký và liên hệ nhân viên; việc gửi lại phụ thuộc cấu hình dịch vụ email của hệ thống.')

h('Chức năng của Quản trị viên',2)
p('Quản trị viên quản lý dữ liệu danh mục, lịch khởi hành, hướng dẫn viên, tài khoản và cấu hình AI; đồng thời có thể thực hiện quy trình tư vấn và kế toán được mô tả ở các mục tiếp theo.')
sub('Quản lý sản phẩm tour')
steps('Mở Quản Lý Tour. Tìm theo mã, tên hoặc điểm đến; sử dụng các bộ lọc trên bảng nếu cần.',
      'Chọn thêm tour và nhập mã tour duy nhất, tên tour, danh mục, điểm đến, số ngày/đêm, giá cơ bản, mô tả, lịch trình và thông tin bổ sung. Không nhập số tiền âm.',
      'Lưu tour, kiểm tra lại trên danh sách và chọn Xem để mở chi tiết đối chiếu. Có thể chọn Mở lịch & HDV ngay trên dòng tour để bổ sung đợt khởi hành.',
      'Khi muốn xóa, kiểm tra đúng tour và đọc xác nhận. Nếu tour đã có đặt chỗ, xóa thường bị chặn và Admin có thể nhận thêm lựa chọn xóa cưỡng chế. Chỉ xác nhận khi đã liên hệ khách, hoàn tiền thực tế và lưu hồ sơ cần thiết: thao tác sẽ xóa tour, lịch khởi hành và dữ liệu đặt chỗ liên quan, không thể hoàn tác từ giao diện.')
figure('04-tours.png','Danh sách sản phẩm tour và thao tác quản trị')
p('Danh sách tour hiện có thao tác xem, mở lịch và xóa; chưa có nút chỉnh sửa nội dung tour trực tiếp. Nếu cần sửa một tour đã lưu, chuyển yêu cầu cho người phụ trách kỹ thuật thực hiện qua chức năng cập nhật được hỗ trợ; không xóa tour đã có giao dịch chỉ để nhập lại nội dung.')
p('Biểu mẫu thêm tour mới có thể tạo mô tả/lịch trình bằng AI ngay trong cửa sổ nhập liệu. Nhập tên hoặc điểm đến, chạy AI rồi đọc và sửa bản nháp trước khi lưu. Giá tour phải lớn hơn 0 và cần nhập đủ các thông tin mà biểu mẫu yêu cầu.')
p('Nếu chọn tạo đợt khởi hành đầu tiên cùng tour, nhập ngày đi, ngày về và tổng số chỗ. Sau khi lưu, kiểm tra cả tour và lịch đi. Hai bước được gửi riêng: khi tour đã tạo nhưng lịch đi lỗi, mở Lịch Khởi Hành & Chỗ để bổ sung, không tạo lại tour với cùng mã.')
figure('15-create-tour.png','Biểu mẫu tạo tour với hỗ trợ AI và lịch khởi hành đầu tiên')
sub('Mở và điều hành đợt khởi hành')
steps('Mở Lịch Khởi Hành & Chỗ, chọn Mở đợt khởi hành mới. Chọn tour, ngày khởi hành, ngày về, tổng số chỗ và giá áp dụng nếu khác giá cơ bản.',
      'Chọn hướng dẫn viên và ghi chú. Kiểm tra thông báo trùng lịch; nếu có xung đột, đổi hướng dẫn viên hoặc điều chỉnh thời gian rồi lưu.',
      'Kiểm tra số chỗ còn và trạng thái trên danh sách. Khi sửa tổng số chỗ, không giảm thấp hơn số chỗ đã đặt. Không thể xóa đợt đã có đặt chỗ.',
      'Khi đủ điều kiện vận hành, chọn Khởi hành và xác nhận. Hệ thống kiểm tra thời điểm, trạng thái chuyến và công nợ; các đơn đã duyệt phải thanh toán đủ.',
      'Sau chuyến, sử dụng thao tác hoàn thành trên đợt đang chạy khi đáp ứng điều kiện. Hệ thống đồng bộ các đặt chỗ liên quan sang hoàn thành. Dùng các nút vòng đời chuyến thay vì chỉ sửa nhãn trạng thái.')
figure('05-departures.png','Lịch khởi hành, số chỗ và phân công hướng dẫn viên')
sub('Xem danh sách khách trong chuyến và hủy chuyến khẩn cấp')
steps('Tại đợt khởi hành cần xử lý, mở danh sách khách. Admin hoặc Tư vấn viên có thể xem mã đơn, họ tên, số điện thoại, số khách, tiền và trạng thái; dùng nút gọi hoặc sao chép số điện thoại để liên hệ.',
      'Khi cần hủy toàn bộ chuyến vì sự cố, Admin mở danh sách khách rồi chọn hủy chuyến khẩn cấp. Nhập rõ lý do trước khi gửi; đây là thao tác ảnh hưởng nhiều đơn cùng lúc.',
      'Kiểm tra kết quả hủy chuyến và các đơn đang chờ/đã duyệt được chuyển sang hủy. Hệ thống ghi nhận hoàn 100% tiền đã đóng cho các đơn thuộc diện xử lý; người vận hành vẫn phải liên hệ khách, thực hiện hoàn tiền bên ngoài và đối soát từng trường hợp.')
figure('16-customers.png','Danh sách khách của đợt khởi hành và thao tác liên hệ')
sub('Xóa vĩnh viễn đặt chỗ')
p('Trong danh sách đặt chỗ, chỉ Admin có quyền xóa vĩnh viễn. Chọn đúng đơn, đọc cảnh báo và chỉ xác nhận khi đã lưu hồ sơ cần thiết, đối soát tiền với khách. Hệ thống xóa cả giao dịch thanh toán liên quan và trả chỗ nếu đơn còn giữ chỗ. Thao tác xóa không thay thế quy trình hủy/hoàn tiền và không tự thực hiện giao dịch ngân hàng.')
sub('Quản lý hồ sơ hướng dẫn viên')
steps('Mở Hướng Dẫn Viên. Tìm theo tên và sử dụng bộ lọc trạng thái/ngôn ngữ để xem hồ sơ cần xử lý.',
      'Chọn thêm mới; nhập họ tên, số điện thoại, email nếu có, trạng thái, ngôn ngữ, kinh nghiệm và các thông tin chuyên môn. Lưu rồi kiểm tra lại hồ sơ.',
      'Dùng nút sửa khi thông tin thay đổi. Việc phân công cho chuyến được thực hiện trên lịch khởi hành; kiểm tra toàn bộ khoảng thời gian trước khi gán.',
      'Trước khi xóa hồ sơ, kiểm tra các chuyến đang được phân công. Phiên bản hiện tại sẽ gỡ liên kết phân công khi xóa; cần bố trí người thay thế trước để tránh thiếu hướng dẫn viên.')
figure('09-guides.png','Danh sách hướng dẫn viên trong dữ liệu demo')
sub('Quản lý tài khoản và phân quyền')
steps('Mở Quản Lý Tài Khoản (RBAC). Chọn thêm tài khoản; nhập tên đăng nhập duy nhất, họ tên, mật khẩu và thông tin liên hệ theo biểu mẫu.',
      'Chọn một vai trò: ADMIN, CONSULTANT hoặc ACCOUNTANT. Cấp đúng quyền theo công việc rồi lưu.',
      'Khi sửa, kiểm tra đúng người dùng trước khi đổi thông tin, mật khẩu, vai trò hoặc trạng thái hoạt động. Thông báo riêng cho người dùng khi có thay đổi tài khoản.',
      'Chỉ khóa hoặc xóa khi đã kiểm tra nhu cầu truy cập. Hệ thống bảo vệ Admin hoạt động cuối cùng; luôn duy trì ít nhất một quản trị viên có thể đăng nhập.')
figure('12-users.png','Quản lý tài khoản nhân viên và vai trò')
sub('Cấu hình AI và tạo nội dung tour')
steps('Đăng nhập Admin, chọn nút trạng thái AI trên thanh trên cùng để mở cấu hình. Nhập khóa Gemini hợp lệ và chọn mô hình khả dụng; lưu rồi kiểm tra trạng thái kết nối. Không chia sẻ ảnh chứa khóa đầy đủ.',
      'Nếu không sử dụng dịch vụ cloud, chọn chế độ xử lý nội bộ. Kết quả nội bộ có thể đơn giản hơn kết quả Gemini.',
      'Mở Trợ lý AI Gemini, chọn phần tạo mô tả tour; nhập tên tour, điểm đến và thông tin biểu mẫu yêu cầu, sau đó chạy tạo mô tả.',
      'Đọc và sửa nội dung AI, kiểm tra lịch trình, dịch vụ và thông tin thực tế. Nếu áp dụng sang biểu mẫu tour, tiếp tục điền các trường bắt buộc rồi lưu. Nội dung AI chỉ là bản nháp, không tự trở thành tour đã lưu.')
figure('13-ai.png','AI Hub hỗ trợ tư vấn và tạo nội dung')

h('Chức năng của Nhân viên tư vấn',2)
p('Tư vấn viên được xử lý yêu cầu khách, tạo đặt chỗ, tra cứu dữ liệu phục vụ tư vấn, xem phản hồi và sử dụng chức năng AI được cấp. Tư vấn viên không quản trị tài khoản, không cấu hình khóa AI và không ghi nhận thu tiền; giao diện kế toán và báo cáo doanh thu dành cho Admin/Kế toán.')
sub('Tiếp nhận và duyệt yêu cầu khách')
steps('Mở Hàng Chờ Tư Vấn. Tìm đúng yêu cầu theo mã, tên khách hoặc thông tin liên hệ; xem tour, đợt khởi hành và số người.',
      'Liên hệ khách để xác minh thông tin và thống nhất lịch, số khách, giá cùng điều kiện đặt tour.',
      'Nếu đủ điều kiện, chọn xác nhận, nhập ghi chú nếu cần rồi nhấn Xác Nhận & Giữ Chỗ. Hệ thống kiểm tra số chỗ còn tại thời điểm duyệt và trừ chỗ khi thành công.',
      'Nếu khách không tiếp tục, chọn từ chối và nêu lý do. Việc từ chối một yêu cầu chưa giữ chỗ không làm tăng số chỗ của chuyến.',
      'Sau duyệt, tìm đơn ở Đặt Chỗ Đã Xác Nhận và bàn giao mã đặt chỗ cho kế toán để thu tiền. Nếu báo không đủ chỗ, đề xuất đợt khác; không xác nhận ngoài hệ thống rồi bỏ qua cảnh báo.')
figure('06-queue.png','Màn hình hàng chờ tiếp nhận yêu cầu khách')
sub('Tạo đặt chỗ trực tiếp và quản lý khách hàng')
steps('Từ tour hoặc lịch khởi hành còn mở, chọn Đặt chỗ. Chọn đúng chuyến, nhập họ tên, số điện thoại, email và số lượng khách.',
      'Đối chiếu thông tin khách đã có để tránh tạo hồ sơ trùng; kiểm tra số tiền tổng trước khi lưu. Với tài khoản tư vấn, không khai đã thu tiền và không nhập một khoản thanh toán thay kế toán.',
      'Nhấn Xác nhận Đặt & Trừ Chỗ. Khi thành công, lưu mã đặt chỗ và kiểm tra số chỗ còn trên lịch khởi hành.',
      'Mở Đặt Chỗ Đã Xác Nhận để tìm đơn, kiểm tra trạng thái, tiền phải thu, tiền đã đóng và ghi chú. Bàn giao công nợ cho kế toán theo mã đơn.')
figure('07-bookings.png','Danh sách đặt chỗ đã tiếp nhận và trạng thái xử lý')
p('Tiền đặt chỗ căn cứ vào giá áp dụng của đợt và số người. Mức cọc mặc định là 30% tổng tiền. Việc giữ chỗ và việc thu tiền là hai bước khác nhau; cần đối chiếu cả trạng thái đặt chỗ lẫn số tiền đã đóng trước khi thông báo cho khách.')
sub('Hủy đặt chỗ theo yêu cầu khách')
steps('Tìm đúng mã đặt chỗ, kiểm tra ngày khởi hành, trạng thái chuyến và số tiền đã đóng. Xác minh yêu cầu hủy với khách.',
      'Chọn hủy, nhập lý do, đọc số tiền hoàn dự kiến rồi xác nhận. Tư vấn viên dùng tỷ lệ mặc định của hệ thống; việc thay đổi tỷ lệ chỉ dành cho Admin hoặc Kế toán.',
      'Kiểm tra trạng thái đã hủy và số chỗ được trả lại nếu đơn từng giữ chỗ. Bàn giao số tiền hoàn cho kế toán để đối soát thực tế; ghi nhận trên phần mềm không đồng nghĩa ngân hàng tự chuyển tiền.',
      'Nếu chuyến đã khởi hành hoặc đã hoàn thành, không tiếp tục hủy theo quy trình trước chuyến; liên hệ quản trị viên để xử lý nghiệp vụ phù hợp.')
p('Chính sách mặc định tính trên tiền đã đóng: hủy trước ít nhất 15 ngày hoàn 100%; từ 7 đến dưới 15 ngày hoàn 70%; từ 3 đến dưới 7 ngày hoàn 50%; dưới 3 ngày hoàn 0%. Kiểm tra số tiền hệ thống tính trước khi xác nhận.')
sub('Theo dõi phản hồi và phân tích AI')
steps('Mở Phản Hồi & Đánh Giá. Tìm theo khách, nội dung hoặc tour; lọc điểm đánh giá và sắc thái phản hồi để xác định vấn đề cần xử lý.',
      'Đọc nhận xét gốc và điểm đánh giá hướng dẫn viên. Chuyển các vấn đề liên quan cho người phụ trách tour.',
      'Trong AI Hub, chọn phân tích phản hồi, chọn tour hoặc phạm vi biểu mẫu hỗ trợ rồi chạy phân tích. Đối chiếu tóm tắt với phản hồi gốc trước khi sử dụng đề xuất cải thiện.')
figure('10-feedbacks.png','Tra cứu phản hồi và đánh giá khách hàng')
sub('Nhập phản hồi mới trên giao diện')
steps('Trong Phản Hồi & Đánh Giá, chọn Gửi Đánh Giá Mới. Chọn tour, nhập họ tên người phản hồi, điểm tour và điểm hướng dẫn viên từ 1 đến 5.',
      'Nhập nhận xét có ít nhất 5 ký tự rồi chọn Gửi Đánh Giá Ngay. Nếu nhập hộ khách, sử dụng đúng nội dung khách cung cấp và kiểm tra lại họ tên; biểu mẫu có thể điền sẵn tên tài khoản đang đăng nhập.',
      'Sau thông báo thành công, kiểm tra phản hồi xuất hiện trên danh sách. Có thể dùng nội dung này trong chức năng phân tích phản hồi của AI Hub; kết quả AI cần được đối chiếu với nhận xét gốc.')
figure('17-feedback-form.png','Biểu mẫu gửi đánh giá mới trong màn hình phản hồi')
p('Không xem kết quả phân tích AI là xác nhận tự động rằng phản ánh đã được giải quyết. Với khách chưa đăng nhập, nhân viên có thể tiếp nhận và nhập phản hồi trên màn hình nghiệp vụ theo thông tin được cung cấp.')

h('Chức năng của Kế toán',2)
p('Kế toán được xem đặt chỗ, ghi nhận thanh toán, xử lý hủy/hoàn theo quyền và xem báo cáo doanh thu. Kế toán không tạo/sửa tour, điều hành lịch khởi hành, duyệt yêu cầu tư vấn hay quản trị tài khoản.')
sub('Ghi nhận tiền cọc và thanh toán còn lại')
steps('Mở Kế Toán & Đặt Cọc hoặc mở đơn từ Đặt Chỗ Đã Xác Nhận. Tìm đúng mã đặt chỗ và kiểm tra khách, tổng tiền, số đã đóng cùng số còn nợ.',
      'Đối chiếu chứng từ hoặc tiền thực nhận trước khi chọn thu tiền. Chỉ ghi nhận cho đơn đã được duyệt; không thu trên yêu cầu còn chờ xác nhận.',
      'Nhập số tiền thu đợt này bằng VNĐ, chọn loại thanh toán và phương thức phù hợp. Số tiền phải lớn hơn 0 và không vượt công nợ còn lại.',
      'Nhấn Xác nhận Thu Tiền. Kiểm tra giao dịch mới, số tiền đã đóng và trạng thái của đơn. Nếu thông báo lỗi hoặc mất kết nối, tra cứu lại trước khi nhập lại khoản thu để tránh ghi trùng.',
      'Khi khách thanh toán đủ, kiểm tra công nợ về 0 và thông báo cho người phụ trách chuyến. Bảo quản chứng từ thực tế theo quy trình đơn vị.')
p('Nút thu tiền trực tiếp trên màn hình kế toán cho phép chọn một đơn còn công nợ, xem thông tin tóm tắt rồi chuyển sang biểu mẫu thu tiền. Nếu danh sách trống, kiểm tra có đơn đã duyệt còn dư nợ hay không.')
figure('08-payments.png','Theo dõi giao dịch thanh toán và công nợ')
p('Hệ thống có hướng dẫn chuyển khoản/VietQR ở phần tra cứu đơn, nhưng chưa có cơ chế đối soát ngân hàng tự động. Các lựa chọn tiền mặt, chuyển khoản hoặc thẻ dùng để ghi nhận phương thức giao dịch. Nhân viên vẫn phải xác minh tiền thực nhận, ghi nhận thủ công và kiểm tra email thông báo khi dịch vụ gửi thư đã được cấu hình.')
sub('Kiểm tra hủy đơn và hoàn tiền')
steps('Mở đúng đặt chỗ cần hủy, đối chiếu số tiền đã thu và chính sách áp dụng. Thực hiện quy trình hủy đã mô tả ở mục tư vấn.',
      'Nếu cần xử lý tỷ lệ hoàn khác mặc định, chỉ thực hiện trong phạm vi quyền Admin/Kế toán và theo quyết định nghiệp vụ được chấp thuận; ghi rõ lý do.',
      'Kiểm tra khoản hoàn được ghi nhận trong sổ giao dịch, thực hiện hoàn tiền thực tế và đối soát chứng từ. Không ghi thêm một giao dịch thu tiền âm thay cho quy trình hoàn tiền.')
sub('Xem báo cáo và đối soát')
steps('Mở Thống Kê Doanh Thu bằng tài khoản Admin hoặc Kế toán. Xem tổng quan và các biểu đồ doanh thu tour, tỷ lệ lấp đầy theo dữ liệu hiện có.',
      'Phân biệt doanh thu dự kiến từ đơn đã duyệt chưa hủy với tiền đã thực thu sau khi trừ tiền hoàn. Đối chiếu mã đơn trong Đặt Chỗ Đã Xác Nhận và giao dịch ở Kế Toán & Đặt Cọc khi có chênh lệch.',
      'Với biểu đồ không có dữ liệu, kiểm tra có đơn đã duyệt/giao dịch hay chưa. Yêu cầu chưa được duyệt không được tính vào doanh thu dự kiến.',
      'Lập báo cáo nghiệp vụ theo dữ liệu đã đối soát. Phiên bản hiện tại chưa có chức năng xuất báo cáo Excel/PDF hoàn chỉnh; không dùng các nút hoặc khả năng chưa được triển khai làm căn cứ hướng dẫn.')
figure('11-analytics.png','Màn hình thống kê doanh thu và vận hành')

h('Xử lý sự cố và lưu ý sử dụng',2)
sub('Không mở được ứng dụng hoặc không tải được dữ liệu')
p('Kiểm tra cửa sổ start.bat còn chạy và đã báo Application startup complete; dùng đúng địa chỉ http://127.0.0.1:8000. Nếu báo thiếu thư viện, cài lại theo requirements-lock.txt trong .venv. Nếu cổng 8000 đang được dùng, nhờ người phụ trách kỹ thuật kiểm tra tiến trình thay vì chạy nhiều máy chủ cùng lúc. Khi lỗi cơ sở dữ liệu, kiểm tra cấu hình kết nối và gửi thông báo lỗi cho quản trị viên.')
sub('Không đăng nhập được hoặc bị từ chối thao tác')
p('Kiểm tra tên đăng nhập, mật khẩu, trạng thái tài khoản và vai trò. Khi phiên hết hạn, đăng nhập lại. Nếu thiếu quyền, nhờ Admin xác nhận phân công; đổi đường dẫn trang không làm thay đổi quyền truy cập. Không chia sẻ tài khoản Admin để khắc phục tạm thời.')
sub('Hết chỗ, trùng lịch hoặc không thể xóa dữ liệu')
p('Tải lại danh sách để xem số chỗ mới nhất. Khi không đủ chỗ, chọn đợt khác hoặc giảm số người sau khi thống nhất với khách. Khi hướng dẫn viên trùng lịch, chọn người khác hoặc sửa khoảng ngày. Nếu xóa thường bị chặn vì có đặt chỗ, kiểm tra lại mục tiêu; lựa chọn xóa cưỡng chế dành cho Admin sẽ làm mất dữ liệu liên quan và chỉ thực hiện sau khi đối soát, lưu hồ sơ cần thiết.')
sub('Không tra cứu được đơn hoặc chưa nhận email')
p('Nhập đúng mã đơn và số điện thoại/email đã đăng ký; không dùng số liên hệ khác với hồ sơ khách. Nếu quên mã, dùng chức năng gửi lại qua email và kiểm tra Spam. Admin kiểm tra SMTP và APP_BASE_URL nếu nhiều khách cùng không nhận được thư hoặc liên kết mở sai địa chỉ. Khi thanh toán chưa cập nhật, chuyển chứng từ cho Kế toán xác minh; không chuyển thêm tiền chỉ vì chưa thấy trạng thái đổi.')
sub('Không thu được tiền hoặc không khởi hành được')
p('Kiểm tra đơn đã được duyệt, số tiền nhập là số dương và không vượt dư nợ. Tư vấn viên cần bàn giao cho Kế toán/Admin để thu tiền. Khi chuyến bị chặn khởi hành/hoàn thành, kiểm tra ngày, trạng thái và tình trạng thanh toán của tất cả đơn đã duyệt; xử lý công nợ trước khi thử lại.')
sub('AI không trả lời hoặc thông tin chưa chính xác')
p('Admin kiểm tra trạng thái AI, khóa, mô hình và kết nối Internet. Có thể chuyển về chế độ nội bộ để tiếp tục thao tác. Nếu tour chưa có phản hồi thì không thể phân tích nội dung chưa tồn tại. Mọi mô tả, gợi ý và tóm tắt AI cần được người dùng kiểm tra; giá, ngày và chỗ phải đối chiếu với dữ liệu tour/chuyến hiện có.')
sub('Bảo quản dữ liệu và liên hệ hỗ trợ')
p('Đăng xuất trên máy dùng chung, không chia sẻ mật khẩu hoặc khóa AI. Quản trị viên sao lưu cơ sở dữ liệu định kỳ và trước khi cập nhật; với SQLite, dừng ứng dụng trước khi sao chép file dữ liệu để bảo đảm bản sao nhất quán. Khi báo lỗi, gửi màn hình chức năng, thời điểm, mã đơn/tour và thông báo lỗi; che mật khẩu, khóa và thông tin cá nhân không cần thiết. Liên hệ người quản trị hoặc nhóm phát triển phụ trách triển khai.')
body.append(section)

# Retain cover controls, fonts, page furniture, numbering and every opaque package part.
title='Hệ thống quản lý tour du lịch có tích hợp AI'
for name in ['word/document.xml','word/glossary/document.xml','docProps/core.xml']:
    tree = root if name=='word/document.xml' else E.fromstring(parts[name])
    for node in tree.iter():
        if node.text and 'Hệ thống quản lý bán hàng có tích hợp AI' in node.text:
            node.text=node.text.replace('Hệ thống quản lý bán hàng có tích hợp AI',title)
    if name!='word/document.xml':parts[name]=ser(tree)

# Rebuild the cached TOC with the retained TOC paragraph styles and live PAGEREF fields.
for node in list(toc):toc.remove(node)
toc.append(toc_title)
counts=[0,0]
for text,level,anchor in entries:
    if level==1:counts[0]+=1;counts[1]=0; prefix=f'{counts[0]}.'
    else:counts[1]+=1;prefix=f'{counts[0]}.{counts[1]}'
    node=clean_p(toc1 if level==1 else toc2,'')
    node.remove(node.find('w:r',NS))
    link=E.SubElement(node,tag('hyperlink'),{tag('anchor'):anchor})
    run=E.SubElement(link,tag('r'))
    rp=E.SubElement(run,tag('rPr'))
    E.SubElement(rp,tag('rFonts'),{tag('ascii'):'Times New Roman',tag('hAnsi'):'Times New Roman'})
    E.SubElement(rp,tag('sz'),{tag('val'):'26'})
    E.SubElement(run,tag('t')).text=f'{prefix} {text}'
    E.SubElement(run,tag('tab'))
    field=E.SubElement(link,tag('fldSimple'),{tag('instr'):f' PAGEREF {anchor} \\h ',tag('dirty'):'true'})
    fr=E.SubElement(field,tag('r'))
    E.SubElement(fr,tag('t')).text='–'
    toc.append(node)
settings=E.fromstring(parts['word/settings.xml'])
update=settings.find('w:updateFields',NS)
if update is None:update=E.SubElement(settings,tag('updateFields'))
update.set(tag('val'),'true')
parts['word/settings.xml']=ser(settings)
parts['word/document.xml']=ser(root)
parts['word/_rels/document.xml.rels']=ser(rels)
types=E.fromstring(parts['[Content_Types].xml'])
if not types.xpath('*[@Extension="png"]'):
    E.SubElement(types,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='png',ContentType='image/png')
parts['[Content_Types].xml']=ser(types)
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
    for name,data in parts.items():z.writestr(name,data)
with ZipFile(REF) as a, ZipFile(OUT) as b:
    for name in a.namelist():
        if name not in ['word/document.xml','word/_rels/document.xml.rels','word/settings.xml','[Content_Types].xml','word/glossary/document.xml','docProps/core.xml']:
            assert a.read(name)==b.read(name), name
    assert E.tostring(E.fromstring(b.read('word/document.xml')).find('.//w:sectPr',NS))==E.tostring(section)
Document(OUT)
report={'reference':str(REF),'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'output':str(OUT),'figures':picture_count,'headings':entries,'preserved':'styles, numbering, margins, header/footer and original reference'}
(QA/'build.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'artifact.md').write_text(f'''# User guide template contract
Reference: {REF}
SHA256: {report['reference_sha256']}
Source: one A4 portrait section; margins left1701/right1134/top1134/bottom1134 twips.
Preserve: cover SDTs, TNR36 title/TNR25 subtitle, blue cover borders, original styles and numbering, all headers/footers and section properties.
Body: retained direct TNR13, black finished text; remove blue italic drafting hints. Heading1/2 cloned from reference; automatic numbering retained.
Slots: replace sales-system title in body/core/glossary; fill intro, hardware/software, actor functions; add guest and troubleshooting groups; refresh cached TOC page references.
New components: 17 actual screenshots from isolated demo server; inline images width5760000EMU with captions; no external imagery or invented UI.
Coverage: body, cover controls, TOC, core properties, glossary, header/footer examined. Original stays unchanged.
QA: preserve-only parts byte-identical; render every page and inspect before delivery.
''',encoding='utf-8')
print(json.dumps({'output':str(OUT),'figures':picture_count,'headings':len(entries)},ensure_ascii=False))
