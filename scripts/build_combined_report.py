"""Synthesize seven project documents into one consistently styled academic report."""
from pathlib import Path
from io import BytesIO
from copy import deepcopy
from zipfile import ZipFile
from collections import Counter
import ast, json, hashlib, re
from lxml import etree as E
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
QA=ROOT/'artifacts/combined-report'
OUT=ROOT/'docs/BAO_CAO_TONG_HOP_VIETJOURNEY_AI.docx'
SOURCES=json.loads((QA/'sources.json').read_text(encoding='utf-8'))
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
DIAGRAMS={int(p.name.split('.')[0]):p for p in (ROOT/'docs/diagrams').glob('*.jpg')}
d=Document()
sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
sec.left_margin=Cm(3);sec.right_margin=Cm(2)
sec.header_distance=Cm(0.8);sec.footer_distance=Cm(0.8)
sec.different_first_page_header_footer=True

def style_font(style,size,bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if 'theme' in key.lower():del fonts.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
for name,size,bold in [('Normal',13,False),('Title',22,True),('Subtitle',16,False),('Heading 1',16,True),('Heading 2',14,True),('Heading 3',13,True),('Caption',12,False),('TOC Heading',16,True)]:
    style_font(d.styles[name],size,bold)
normal=d.styles['Normal'].paragraph_format
normal.line_spacing=1.15;normal.space_after=Pt(6)
normal.widow_control=True
for name in ['Heading 1','Heading 2','Heading 3']:
    pf=d.styles[name].paragraph_format;pf.space_before=Pt(10);pf.space_after=Pt(6);pf.keep_with_next=True;pf.keep_together=True
d.styles['Caption'].paragraph_format.space_after=Pt(8)
d.styles['Caption'].paragraph_format.line_spacing=1.0
d.core_properties.title='Báo cáo tổng hợp hệ thống quản lý tour du lịch có tích hợp AI'
d.core_properties.subject='VietJourney AI'
d.core_properties.author='Nhóm 17 – Hà Thành Duy, Bùi Văn Trường'
d.core_properties.keywords='VietJourney AI, tour du lịch, phân tích thiết kế, kiểm thử, hướng dẫn sử dụng'
header=sec.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=header.add_run('VIETJOURNEY AI | NHÓM 17');run.font.name='Times New Roman';run.font.size=Pt(10)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' PAGE ')
fr=OxmlElement('w:r');ft=OxmlElement('w:t');ft.text='1';fr.append(ft);f.append(fr);footer._p.append(f)
settings=d.settings.element
up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');settings.append(up)

headings=[];figures=[];tables=[]
def p(text,bold=False,center=False):
    x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    x.add_run(text).bold=bold
    return x
def h(text,level=1,page=False,toc=True):
    x=d.add_paragraph(text,style=f'Heading {level}')
    if page:x.paragraph_format.page_break_before=True
    bookmark=f'ReportHeading{len(headings)+100}'
    bid=str(len(headings)+100)
    st=OxmlElement('w:bookmarkStart');st.set(qn('w:id'),bid);st.set(qn('w:name'),bookmark)
    en=OxmlElement('w:bookmarkEnd');en.set(qn('w:id'),bid)
    x._p.insert(1 if x._p.pPr is not None else 0,st);x._p.append(en)
    headings.append({'text':text,'level':level,'anchor':bookmark,'toc':toc})
    return x
def sub(text):
    x=p(text,True);x.paragraph_format.keep_with_next=True
    return x
def steps(items):
    for i,t in enumerate(items,1):p(f'{i}. {t}')
def note(text):
    x=p(text)
    for r in x.runs:r.italic=True;r.font.size=Pt(11)
    return x
def caption(text):
    x=d.add_paragraph(text,style='Caption');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    x.paragraph_format.keep_together=True
    return x
def table(headers,rows,widths,title,size=11):
    tables.append(title)
    cap=caption(f'Bảng {len(tables)}. {title}');cap.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for j,txt in enumerate(headers):t.rows[0].cells[j].text=str(txt)
    for row in rows:
        cells=t.add_row().cells
        for j,txt in enumerate(row):cells[j].text=str(txt)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');pr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6' if ri==0 else ('F7F7F7' if ri%2==0 else 'FFFFFF'));cp.append(shade)
            for pp in cell.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1.05
                pp.paragraph_format.widow_control=True
                for rr in pp.runs:rr.font.name='Times New Roman';rr.font.size=Pt(size);rr.bold=ri==0
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
    t._tbl.tblPr.append(borders)
    return t
def figure(data,title,width=16,height=20,source=None):
    blob=Path(data).read_bytes() if isinstance(data,(str,Path)) else data
    with Image.open(BytesIO(blob)) as im:iw,ih=im.size
    scale=min(width/iw,height/ih)
    pp=d.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.keep_with_next=True;pp.paragraph_format.space_after=Pt(3)
    pp.add_run().add_picture(BytesIO(blob),width=Cm(iw*scale),height=Cm(ih*scale))
    figures.append({'number':len(figures)+1,'title':title,'source':source,'sha256':hashlib.sha256(blob).hexdigest()})
    caption(f'Hình {len(figures)}. {title}')

# Cover uses the same typeface and A4 page system as the project documents.
for _ in range(3):p('')
p('BÁO CÁO TỔNG HỢP DỰ ÁN',True,True).runs[0].font.size=Pt(18)
p('HỆ THỐNG QUẢN LÝ TOUR DU LỊCH CÓ TÍCH HỢP AI',True,True).runs[0].font.size=Pt(22)
p('VIETJOURNEY AI',True,True).runs[0].font.size=Pt(20)
p('Phân tích yêu cầu, thiết kế, phát triển, kiểm thử và vận hành',False,True)
for _ in range(2):p('')
p('Nhóm 17',True,True)
p('Hà Thành Duy',False,True);p('Bùi Văn Trường',False,True)
p('Kế hoạch thực hiện: 27/07/2026 – 27/09/2026',False,True)
p('Bản tổng hợp cập nhật ngày 29/10/2026',False,True)
d.add_page_break()
toc_title=p('MỤC LỤC',True,True);toc_title.runs[0].font.size=Pt(16)
toc_marker=d.add_paragraph('')

h('TÓM TẮT BÁO CÁO',page=True)
p('VietJourney AI là hệ thống quản lý tour du lịch kết hợp các chức năng điều hành, bán hàng, kế toán và hỗ trợ AI. Báo cáo tổng hợp hồ sơ của Nhóm 17 thành một tài liệu xuyên suốt từ kế hoạch thực hiện đến yêu cầu, mô hình thiết kế, dữ liệu, kiểm thử và hướng dẫn sử dụng. Hệ thống phục vụ khách vãng lai cùng ba vai trò nội bộ: quản trị viên, nhân viên tư vấn và kế toán.')
p('Quy trình trọng tâm gồm tiếp nhận nhu cầu du lịch, kiểm tra lịch và số chỗ, duyệt giữ chỗ, thu tiền, điều hành chuyến, ghi nhận phản hồi và theo dõi báo cáo. AI hỗ trợ ba việc: tư vấn trên danh sách tour thực có, tạo bản nháp mô tả/lịch trình và tổng hợp nhận xét. Người vận hành kiểm tra nội dung AI trước khi sử dụng; việc giữ chỗ và ghi nhận tiền vẫn tuân theo quy tắc nghiệp vụ của hệ thống.')
p('Bản báo cáo sử dụng bảy tài liệu nội bộ đã hoàn thiện, trong đó hướng dẫn sử dụng cập nhật ngày 29/10/2026 phản ánh các bổ sung về tra cứu đơn, email, VietQR và thao tác quản trị. Các nội dung lặp được biên tập lại; sơ đồ do nhóm cung cấp được giữ nguyên. Những khác biệt giữa yêu cầu ban đầu, thiết kế khái niệm và giao diện hiện tại được giải thích tại phần liên quan.')
p('Hồ sơ kiểm thử gồm 49 tình huống, được ghi ở hai phần: mô tả tình huống và kết quả lịch sử ngày 25/9. Hai phần này không được tính thành 98 ca độc lập. Bản tổng hợp giữ kết quả gốc, đánh dấu nội dung cần kiểm thử lại và không sử dụng ảnh minh họa để thay thế bằng chứng nghiệm thu chức năng.')
table(['Tài liệu nguồn','Nội dung được tổng hợp'],[(s['id'],desc) for s,desc in zip(SOURCES,['Kế hoạch chín tuần, nhân sự và phân công.','Câu hỏi làm rõ, yêu cầu chức năng và phi chức năng.','Tác nhân, tám use case và các biểu đồ hành vi.','Mô hình lớp và đặc tả lớp/dịch vụ.','Môi trường, tình huống và kết quả kiểm thử lịch sử.','Screen Flow, thiết kế dữ liệu và các ràng buộc.','Cài đặt, hướng dẫn thao tác, ảnh giao diện và xử lý sự cố.'])],[2.5,13.5],'Phạm vi sử dụng bảy tài liệu nội bộ',12)

h('CHƯƠNG 1 TỔNG QUAN VÀ KẾ HOẠCH THỰC HIỆN',page=True)
h('1.1 Bài toán và mục tiêu',2)
p('Hoạt động quản lý tour cần phối hợp dữ liệu chương trình du lịch, từng đợt khởi hành, nhân sự hướng dẫn, thông tin khách và các khoản thu. Khi các nghiệp vụ này tách rời, nhân viên khó theo dõi số chỗ và công nợ nhất quán. Bài toán đặt ra là tổ chức dữ liệu tập trung, phân định quyền thao tác và hỗ trợ tư vấn trên lịch mở bán thực tế.')
p('Mục tiêu của hệ thống là quản lý trọn quy trình từ yêu cầu khách đến kết thúc chuyến; kiểm soát số chỗ trước khi nhận đơn; theo dõi cọc, thanh toán và hoàn tiền; cung cấp dữ liệu doanh thu và chất lượng dịch vụ. AI được tích hợp vào công việc tư vấn và biên soạn nội dung để giảm thao tác nhập liệu, đồng thời duy trì bước kiểm tra của con người. [TL02, TL03]')
h('1.2 Phạm vi và đối tượng sử dụng',2)
p('Phạm vi nghiệp vụ gồm tour, lịch khởi hành, hướng dẫn viên, khách hàng và đặt chỗ, thanh toán, phản hồi, tài khoản và báo cáo. Khách vãng lai sử dụng các chức năng công khai; nhân viên đăng nhập để thực hiện nghiệp vụ theo vai trò. Hồ sơ hướng dẫn viên là dữ liệu điều hành, chưa phải một nhóm tài khoản có cổng sử dụng riêng.')
p('Ứng dụng vận hành theo mô hình trình duyệt và máy chủ FastAPI. Giao diện HTML/JavaScript sử dụng API của cùng hệ thống; SQLAlchemy quản lý dữ liệu quan hệ. SQLite được dùng cho môi trường cục bộ, còn MySQL là lựa chọn được mã nguồn hỗ trợ và cần kiểm thử riêng trên môi trường triển khai. [TL03, TL06, TL07]')
h('1.3 Nhóm thực hiện và kế hoạch chín tuần',2)
p('Nhóm 17 gồm Hà Thành Duy và Bùi Văn Trường. Kế hoạch gốc kéo dài từ 27/07/2026 đến 27/09/2026; ngày tổng hợp báo cáo là ngày cập nhật hồ sơ, không thay thế thời gian của kế hoạch ban đầu. [TL01]')
plan_groups={}
for row in SOURCES[0]['tables'][0][1:]:
    if row[1].strip():
        dates=re.findall(r'\d{2}/\d{2}/\d{4}',row[0])
        week=re.search(r'Tuần\s+(\d+)',row[0]).group(1)
        period=f'Tuần {week}\n{dates[0][:5]}–{dates[1][:5]}'
        plan_groups.setdefault(period,[]).append((' '.join(row[1].split()),row[2].strip()))
plan=[]
for period,items in plan_groups.items():
    plan.append((period,'\n'.join(f'{j+1}. {x[0]}' for j,x in enumerate(items)),'\n'.join(f'{j+1}. {x[1]}' for j,x in enumerate(items))))
p('Kế hoạch đi từ khảo sát và mô hình hóa đến xây dựng dữ liệu/API, phát triển nghiệp vụ, tích hợp AI, kiểm thử và hoàn thiện hồ sơ. Các ngày trong bảng dưới thuộc năm 2026; chỉ các công việc được giao cụ thể trong nguồn được liệt kê.')
table(['Tuần và ngày','Công việc','Phụ trách'],plan,[2.8,9.2,4],'Phân công công việc theo kế hoạch',11)

h('CHƯƠNG 2 KHẢO SÁT VÀ ĐẶC TẢ YÊU CẦU',page=True)
h('2.1 Kết quả làm rõ yêu cầu',2)
p('Bộ câu hỏi khảo sát xác định nhu cầu quản lý tập trung dữ liệu lữ hành, kiểm tra lịch hướng dẫn viên, ngăn đặt vượt chỗ, quản lý tiền và sử dụng AI trên dữ liệu có thật. Yêu cầu về giá, ngày đi và số chỗ phải được lấy từ hệ thống; nội dung do mô hình ngôn ngữ tạo ra không được tự thay thế dữ liệu giao dịch. [TL02]')
table(['Nhóm yêu cầu','Nội dung chính'],[
('Dữ liệu tour và chuyến','Tour lưu mã, tên, điểm đến, thời lượng và giá cơ bản; mỗi tour có nhiều đợt với ngày đi/về, sức chứa, giá áp dụng và hướng dẫn viên.'),
('Tiếp nhận và đặt chỗ','Khách gửi yêu cầu chờ duyệt; nhân viên liên hệ, kiểm tra và giữ chỗ khi duyệt hoặc tạo đơn trực tiếp.'),
('Thanh toán và hủy','Theo dõi cọc và tiền còn nợ; ghi nhận giao dịch thu; khi hủy tính tiền hoàn và trả chỗ cho đơn đã giữ chỗ.'),
('Nhân sự và chất lượng','Quản lý hồ sơ hướng dẫn viên, cảnh báo lịch trùng, thu thập phản hồi và phân tích chất lượng.'),
('AI','Tư vấn theo nhu cầu, ngân sách, số ngày; tạo mô tả/lịch trình; tổng hợp phản hồi theo tour.'),
('Quản trị và báo cáo','Đăng nhập, phân quyền, quản lý tài khoản và xem các chỉ số doanh thu, tour bán tốt, tỷ lệ lấp đầy.')],[3.7,12.3],'Tổng hợp các nhóm yêu cầu từ khảo sát',12)
h('2.2 Tác nhân và phạm vi quyền',2)
table(['Tác nhân','Phạm vi thao tác'],[
('Khách vãng lai','Xem tour/hồ sơ hướng dẫn viên, hỏi AI, gửi yêu cầu, tra cứu đơn bằng mã và thông tin liên hệ; sử dụng liên kết email khi được cung cấp.'),
('ADMIN','Quản trị tour, chuyến, hướng dẫn viên, tài khoản, cấu hình AI; thực hiện nghiệp vụ tư vấn/kế toán và các thao tác quản trị đặc biệt.'),
('CONSULTANT','Tư vấn, xử lý hàng chờ, tạo đặt chỗ, theo dõi đơn và nhập/xem phản hồi theo giao diện; không ghi nhận thu tiền hoặc quản trị tài khoản.'),
('ACCOUNTANT','Theo dõi đơn và công nợ, ghi nhận thu tiền, xử lý hủy/hoàn theo quyền, xem dữ liệu doanh thu.'),
('Dịch vụ AI','Thành phần tích hợp xử lý yêu cầu tư vấn, sinh nội dung và tổng hợp văn bản; không tự cấp quyền hoặc quyết định giao dịch.')],[3.5,12.5],'Tác nhân và phạm vi sử dụng',12)
p('Tài liệu khảo sát dùng tên Advisor cho nhân viên tư vấn; báo cáo thống nhất với mã vai trò CONSULTANT. Khách vãng lai là đối tượng sử dụng công khai bổ sung từ luồng vận hành và tài liệu hướng dẫn; hệ thống AI là thành phần tích hợp, không phải tài khoản nhân viên. [TL02, TL03, TL07]')
figure(DIAGRAMS[1],'Use Case tổng quát',height=18,source='TL03')
h('2.3 Phân cấp chức năng',2)
figure(DIAGRAMS[2],'Phân cấp chức năng hệ thống',height=17,source='TL02')
p('Các chức năng nghiệp vụ được tổ chức theo đối tượng dữ liệu và vai trò. Ba chức năng AI nằm trong cùng nền tảng, sử dụng dữ liệu tour/phản hồi làm ngữ cảnh và trả kết quả về giao diện để người dùng xem xét. Việc ẩn nút trên giao diện hỗ trợ trải nghiệm; kiểm tra quyền ở API là lớp quyết định cho thao tác nội bộ.')
h('2.4 Quy tắc nghiệp vụ trọng tâm',2)
table(['Mã','Quy tắc'],[
('BR01','Yêu cầu WAITING_CONFIRM chưa giữ chỗ; chỉ duyệt hoặc đặt trực tiếp hợp lệ mới trừ chỗ.'),
('BR02','Số người là số nguyên dương; kiểm tra đủ chỗ và trạng thái/ngày của chuyến trước khi nhận đặt.'),
('BR03','Giá trị đơn bằng giá áp dụng của đợt nhân số người; giá riêng của đợt được ưu tiên khi có.'),
('BR04','Mức cọc mặc định 30%; dưới ngưỡng vẫn PENDING, đủ cọc là DEPOSITED, đủ tổng tiền là CONFIRMED.'),
('BR05','Khoản thu phải dương và không vượt dư nợ; CONSULTANT không được khai thu tiền.'),
('BR06','Hủy thông thường trước khởi hành: từ 15 ngày hoàn 100%, từ 7 ngày hoàn 70%, từ 3 ngày hoàn 50%, dưới 3 ngày hoàn 0% tiền đã đóng.'),
('BR07','Chỉ trả ghế cho đơn từng giữ chỗ; xử lý hủy phải tránh ghi nhận hoàn tiền lặp.'),
('BR08','Khoảng lịch hướng dẫn viên không được xung đột; giảm sức chứa không thấp hơn số chỗ đã nhận.'),
('BR09','Doanh thu dự kiến từ đơn đã duyệt chưa hủy; tiền thu ròng dựa trên giao dịch thu trừ hoàn.'),
('BR10','Hủy chuyến khẩn cấp và xóa cưỡng chế là thao tác Admin riêng, khác với chính sách hủy đơn thông thường; phải đối soát trước khi xóa dữ liệu.')],[1.8,14.2],'Quy tắc nghiệp vụ tổng hợp',11.5)
p('BR10 phản ánh phần cập nhật trong hướng dẫn sử dụng: hệ thống có thao tác hủy chuyến khẩn cấp ghi nhận hoàn 100% và thao tác xóa dữ liệu dành riêng cho Admin. Các yêu cầu cũ về chặn xóa tour có booking vẫn mô tả luồng xóa thường, không được hiểu là đã loại bỏ mọi quyền xóa quản trị. [TL02, TL03, TL07]')
h('2.5 Đặc tả tám use case chính',2)
ucs=[
('UC001','Đăng nhập và phân quyền','ADMIN, CONSULTANT, ACCOUNTANT','Tài khoản tồn tại, còn hoạt động và được cấp vai trò.','Người dùng nhập tài khoản → hệ thống xác minh mật khẩu băm → cấp token → giao diện hiển thị quyền tương ứng.','Sai thông tin hoặc tài khoản không hợp lệ: từ chối đăng nhập. Đăng xuất xóa thông tin phiên cục bộ; không giải mã mật khẩu.'),
('UC002','Quản lý tour và AI sinh nội dung','ADMIN; dịch vụ AI hỗ trợ','Đăng nhập Admin; có thông tin tên, điểm đến, giá và thời lượng cần nhập.','Xem danh sách → thêm tour → có thể tạo mô tả/lịch trình bằng AI → kiểm tra bản nháp → lưu tour; có thể tạo tiếp lịch đi đầu tiên.','Thiếu dữ liệu: báo lỗi. AI cloud không khả dụng: dùng chế độ nội bộ. Giao diện hiện có xem, mở lịch và xóa; API cập nhật tour là khả năng kỹ thuật, chưa có nút sửa trực tiếp trong bảng tour.'),
('UC003','Quản lý lịch khởi hành','ADMIN','Tour đã tồn tại; có ngày đi/về và sức chứa hợp lệ.','Chọn tour → nhập lịch, số chỗ, giá áp dụng → chọn hướng dẫn viên → kiểm tra trùng lịch → lưu. Khi vận hành, dùng các thao tác khởi hành và hoàn thành.','Không giảm tổng chỗ dưới số đã nhận; xử lý thông báo xung đột hướng dẫn viên. Xóa thường bị chặn khi có booking; quản trị đặc biệt được mô tả ở phần vận hành.'),
('UC004','Khách hàng và đặt chỗ','Khách vãng lai, ADMIN, CONSULTANT','Có lịch mở bán còn chỗ; nhân viên đăng nhập khi tạo/duyệt đơn.','Khách gửi yêu cầu hoặc nhân viên đặt trực tiếp → nhận thông tin khách → kiểm tra và trừ chỗ trong giao dịch khi duyệt/đặt → lưu booking và trả mã.','Hết chỗ hoặc chuyến không hợp lệ: không giữ chỗ. Yêu cầu chưa duyệt được lưu riêng về trạng thái; tra cứu đơn công khai cần mã và thông tin liên hệ khớp.'),
('UC005','Thanh toán và hoàn tiền','ADMIN, ACCOUNTANT','Có đơn đã duyệt còn dư nợ; nhân viên đã xác minh khoản tiền thực nhận.','Tra cứu đơn → xem tổng/đã đóng/còn nợ → nhập tiền và phương thức → lưu giao dịch → cập nhật trạng thái. Khi hủy, tính khoản hoàn theo chính sách.','Không nhận số tiền âm, bằng 0 hoặc vượt dư nợ; tư vấn viên không thu tiền. VietQR hướng dẫn chuyển khoản, chưa đối soát ngân hàng tự động.'),
('UC006','Phản hồi và AI tóm tắt','ADMIN, CONSULTANT; dịch vụ AI','Có tour để gắn phản hồi; muốn phân tích cần có nhận xét.','Nhập/xem đánh giá → lưu điểm và nội dung → chọn tour trong AI Hub → tổng hợp nhận xét → hiển thị điểm nổi bật, vấn đề và đề xuất.','Điểm từ 1–5. Giao diện yêu cầu nhận xét tối thiểu 5 ký tự. Nghiệp vụ hướng đến khách sau chuyến nhưng chưa xác minh bắt buộc người gửi đã hoàn thành booking.'),
('UC007','Quản lý hướng dẫn viên','ADMIN; các đối tượng khác xem theo giao diện','Admin đăng nhập khi tạo/sửa/xóa; có thông tin hồ sơ hợp lệ.','Xem hồ sơ → thêm/cập nhật tên, liên hệ, ngôn ngữ, kinh nghiệm → lưu → kiểm tra lịch trước khi phân công.','Model chưa có trường thẻ HDV hoặc UNIQUE số điện thoại HDV. Xóa hồ sơ có thể gỡ liên kết phân công, cần bố trí người thay thế.'),
('UC008','Thống kê và báo cáo','ADMIN, ACCOUNTANT','Có quyền tài chính và dữ liệu nghiệp vụ.','Mở thống kê → hệ thống tổng hợp doanh thu/tour và tỷ lệ lấp đầy → hiển thị biểu đồ → đối chiếu đặt chỗ và sổ giao dịch.','Yêu cầu chưa duyệt không tính doanh thu dự kiến. Lọc thời gian và xuất Excel/PDF là phạm vi chưa hoàn thiện trong phiên bản mô tả.')]
for code,title,actors,pre,flow,alt in ucs:
    sub(f'{code} {title}')
    p(f'Tác nhân và điều kiện: {actors}. {pre}')
    p(f'Luồng chính: {flow}')
    p(f'Ngoại lệ và giới hạn: {alt}')
p('Biểu đồ Activity và Sequence của UC001–UC008 được tập hợp tại Phụ lục A để phần đặc tả có thể đọc liền mạch. Các luồng được biên tập từ SRS, đối chiếu với hướng dẫn sử dụng cập nhật. [TL03, TL07]')
h('2.6 Yêu cầu phi chức năng và tiêu chí đánh giá',2)
table(['Nhóm','Yêu cầu và cách đánh giá'],[
('Hiệu năng','Đo thời gian CRUD/tìm kiếm và AI riêng trên môi trường xác định. Hồ sơ chưa có số đo hoặc ngưỡng nghiệm thu được xác nhận.'),
('Bảo mật','Xác thực tài khoản, kiểm tra quyền tại API, bảo vệ dữ liệu đơn và tài chính; không lộ mật khẩu, secret hoặc khóa AI.'),
('Chính xác','Giá/ngày/chỗ từ cơ sở dữ liệu; tiền xử lý nguyên VNĐ; kiểm tra bất biến số chỗ và công nợ.'),
('Tin cậy','Nghiệp vụ liên quan được lưu nhất quán; lỗi phải rollback; có sao lưu và kiểm thử gửi lặp, đồng thời.'),
('Dễ sử dụng','Nhãn trạng thái rõ, biểu mẫu hợp lệ, thông báo lỗi có thể xử lý; đánh giá trên màn hình máy tính và màn hình nhỏ.'),
('Bảo trì và mở rộng','Tách giao diện, API, CRUD, mô hình và AI; có kiểm thử và tài liệu vận hành. Thay đổi phải được kiểm thử lại.')],[3,13],'Yêu cầu phi chức năng và định hướng kiểm chứng',12)

h('CHƯƠNG 3 KIẾN TRÚC VÀ THIẾT KẾ HỆ THỐNG',page=True)
h('3.1 Kiến trúc và công nghệ',2)
p('Ứng dụng tổ chức theo luồng trình duyệt → API FastAPI → nghiệp vụ CRUD/dịch vụ AI → cơ sở dữ liệu hoặc nhà cung cấp AI. HTML, JavaScript và Tailwind CSS xây dựng giao diện; Chart.js trình bày số liệu; SQLAlchemy ánh xạ dữ liệu quan hệ; Pydantic kiểm tra cấu trúc dữ liệu đầu vào/đầu ra. Thiết kế phân lớp giúp tập trung kiểm tra quyền, dữ liệu và nghiệp vụ thay vì chỉ phụ thuộc các nút hiển thị. [TL03, TL04, TL06, TL07]')
table(['Lớp','Trách nhiệm','Thành phần đối chiếu'],[
('Giao diện','Điều hướng, biểu mẫu, trạng thái tải, ảnh và biểu đồ; hiển thị kết quả API.','frontend/*.html; frontend/js/app.js; ai_hub.js'),
('API','Tiếp nhận HTTP, xác thực/phân quyền và ánh xạ dữ liệu.','backend/app/routers; schemas.py; auth.py'),
('Nghiệp vụ','Đặt/hủy/duyệt, kiểm tra tiền và số chỗ, tổng hợp số liệu.','backend/app/crud.py'),
('Lưu trữ','Mô hình dữ liệu, quan hệ và kết nối SQLite/MySQL.','models.py; database.py; money.py'),
('AI và email','Tạo ngữ cảnh, gọi mô hình, chuẩn hóa kết quả; gửi thông báo khi có SMTP.','ai_service.py; ai_router.py; email_service.py')],[2.6,8,5.4],'Phân chia trách nhiệm các thành phần',11)
h('3.2 Luồng màn hình',2)
figure(DIAGRAMS[5],'Screen Flow của ứng dụng',height=18,source='TL06')
p('Khách bắt đầu tại trang chủ, có thể xem tour, hỏi AI, gửi yêu cầu và tra cứu đơn. Sau đăng nhập, nhân viên truy cập các màn hình quản lý phù hợp. Quy trình tư vấn đi qua hàng chờ và đặt chỗ; quy trình tài chính đi qua sổ giao dịch và thống kê. Mã đặt chỗ là thông tin xuyên suốt để liên hệ, thanh toán, tra cứu và đối soát. [TL06, TL07]')
h('3.3 Mô hình lớp',2)
figure(DIAGRAMS[3],'Mô hình lớp do nhóm cung cấp',height=21,source='TL04')
p('Mô hình lớp mô tả các thực thể Tour, Departure, Booking, Customer, Payment, TourGuide, Feedback và các thành phần hỗ trợ. GeminiAIService thể hiện trách nhiệm dịch vụ ở mức thiết kế; trong triển khai hiện tại, trách nhiệm này được phân phối trong các hàm của ai_service.py và router. Không coi tên phương thức khái niệm là tên hàm bắt buộc có trong mã. [TL04]')
h('3.4 Ánh xạ trách nhiệm và phương thức',2)
table(['Thành phần thiết kế','Dữ liệu/trách nhiệm','Luồng xử lý'],[
('Tour / generateDescriptionWithAI','Mã, tên, điểm đến, giá cơ bản, ngày/đêm, mô tả.','Gửi tên/điểm đến/thời lượng vào dịch vụ AI; trả bản nháp để Admin duyệt và lưu.'),
('Departure / updateRemainingSeats','Tour, ngày đi/về, tổng chỗ, chỗ còn, HDV.','Kiểm tra đủ chỗ và cập nhật có điều kiện trong giao dịch; cộng lại chỗ khi hủy hợp lệ.'),
('Booking / updatePaidStatus','Khách, chuyến, số người, tổng tiền, đã thu và trạng thái.','Cộng tiền hợp lệ; so ngưỡng cọc 30% và tổng tiền để cập nhật trạng thái.'),
('GeminiAIService / consultTour','Nhu cầu khách, tour còn chỗ, cấu hình mô hình.','Tạo ngữ cảnh, gọi AI, lọc ID hợp lệ và dựng dữ liệu hiển thị từ database.')],[4.3,5.7,6],'Ánh xạ mô hình lớp sang xử lý nghiệp vụ',11)
h('3.5 Bảo đảm nhất quán số chỗ và giao dịch',2)
p('Khi duyệt yêu cầu hoặc tạo booking trực tiếp, nghiệp vụ kiểm tra số người, trạng thái chuyến và ngày đi. Cập nhật chỗ được thực hiện bằng UPDATE có điều kiện đủ chỗ; kết quả không cập nhật được bản ghi sẽ được xem là thất bại. Các thay đổi khách hàng, booking và giao dịch liên quan cần cùng ranh giới transaction để rollback khi có lỗi. SQLite sử dụng BEGIN IMMEDIATE ở các luồng cần đồng bộ; MySQL có khóa bản ghi ở nghiệp vụ liên quan. [TL02, TL03, TL04]')
p('booking_request_keys hỗ trợ nhận diện yêu cầu đã gửi bằng khóa, dấu vân tay nội dung và booking liên quan. Thanh toán kiểm tra mã giao dịch trong phạm vi đơn để hạn chế ghi nhận trùng. Các cơ chế này là thiết kế bảo vệ tính nhất quán; mức bảo đảm thực tế cần được xác nhận lại khi thay đổi backend dữ liệu hoặc các luồng quản trị mới.')

h('CHƯƠNG 4 THIẾT KẾ CƠ SỞ DỮ LIỆU',page=True)
h('4.1 Mô hình dữ liệu quan hệ',2)
p('Thiết kế dữ liệu tổ chức thành các nhóm quản trị, tour và lịch khởi hành, khách hàng và đặt chỗ, tài chính và phản hồi, cùng nhật ký AI và bảng chống gửi lặp. Tài liệu TL06 mô tả bảy bảng chính và tổng quan mười bảng. Danh mục bên dưới bổ sung đủ mười bảng theo mô hình mã nguồn đối chiếu để báo cáo có thể tra cứu thống nhất. [TL06, MK01]')
figure(DIAGRAMS[4],'Sơ đồ ERD do nhóm cung cấp',height=20,source='TL06')
table(['Quan hệ','Ý nghĩa'],[
('tours → departures (1:N)','Một tour có nhiều đợt khởi hành.'),('tour_guides → departures (1:N)','Một hướng dẫn viên có nhiều lịch ở các khoảng thời gian không xung đột.'),
('customers → bookings (1:N)','Một khách có nhiều lần đặt tour.'),('departures → bookings (1:N)','Một chuyến nhận nhiều đơn, trong giới hạn số chỗ.'),
('bookings → payments (1:N)','Một đơn có nhiều khoản thu hoặc hoàn.'),('tours → feedbacks (1:N)','Một tour có nhiều phản hồi.'),
('bookings → feedbacks (1:N, tùy chọn)','Phản hồi có thể gắn với booking; booking_id có thể để trống.'),('bookings → booking_request_keys','Khóa yêu cầu liên kết booking đã được tạo.')],[6,10],'Quan hệ dữ liệu chủ yếu',12)
h('4.2 Từ điển dữ liệu',2)
p('PK là khóa chính; FK là khóa ngoại; UK là ràng buộc duy nhất. Kiểu, NULL và khóa dưới đây đối chiếu với models.py; các giá trị mặc định của ORM không đồng nghĩa mặc định server của DBMS. Các trường Money được biểu diễn bằng DECIMAL(18,0) ở lớp lưu trữ, sử dụng số nguyên VNĐ. [TL06, MK01]')
descriptions={
'users':'Tài khoản nhân sự, xác thực và vai trò.', 'tours':'Chương trình tour, điểm đến, giá và nội dung giới thiệu.',
'tour_guides':'Hồ sơ hướng dẫn viên và trạng thái công tác.', 'departures':'Từng chuyến khởi hành, sức chứa, chỗ còn và phân công.',
'customers':'Thông tin người đặt; số điện thoại định danh duy nhất.', 'bookings':'Đơn đặt tour, số khách, tiền và trạng thái xử lý.',
'payments':'Sổ giao dịch thu/hoàn gắn với đơn.', 'feedbacks':'Đánh giá tour, hướng dẫn viên và nội dung nhận xét.',
'ai_session_logs':'Nhật ký đầu vào, đầu ra và mô hình AI; chưa có khóa ngoại user_id.', 'booking_request_keys':'Khóa chống gửi lại yêu cầu tạo đơn.'}
field_desc={
'id':'Định danh tự tăng','username':'Tên đăng nhập','hashed_password':'Mật khẩu đã băm','full_name':'Họ tên','email':'Địa chỉ email','phone':'Số điện thoại','role':'Vai trò nội bộ','is_active':'Trạng thái hoạt động','created_at':'Thời điểm tạo','updated_at':'Thời điểm cập nhật','code':'Mã tour','title':'Tên tour','destination':'Điểm đến','category':'Danh mục tour','duration_days':'Số ngày','duration_nights':'Số đêm','base_price':'Giá cơ bản/người','description':'Mô tả','itinerary_summary':'Lịch trình tóm tắt','highlights':'Điểm nổi bật','image_url':'Địa chỉ ảnh','languages':'Ngôn ngữ','experience_years':'Số năm kinh nghiệm','status':'Trạng thái','rating':'Điểm đánh giá','notes':'Ghi chú','tour_id':'Tour liên quan','departure_date':'Ngày đi','return_date':'Ngày về','total_slots':'Tổng sức chứa','available_slots':'Chỗ còn','price_override':'Giá riêng của đợt','guide_id':'Hướng dẫn viên được gán','id_card':'Giấy tờ tùy chọn','address':'Địa chỉ','booking_code':'Mã đơn đặt tour','departure_id':'Đợt khởi hành','customer_id':'Khách đặt','num_people':'Số khách','total_amount':'Tổng giá trị đơn','deposit_amount':'Mức cọc yêu cầu','paid_amount':'Tổng tiền đã đóng','refund_amount':'Tiền hoàn','cancellation_reason':'Lý do hủy','special_requests':'Yêu cầu đặc biệt','access_token':'Mã truy cập tra cứu từ email','booking_id':'Đơn liên quan','payment_type':'Loại thu/hoàn','amount':'Số tiền giao dịch','payment_method':'Phương thức','transaction_code':'Mã đối soát','note':'Diễn giải giao dịch','created_by':'Người lập dạng chuỗi','customer_name':'Tên người phản hồi','comments':'Nhận xét','guide_rating':'Điểm hướng dẫn viên','sentiment':'Sắc thái đánh giá','feature_type':'Nhóm chức năng AI','input_prompt':'Đầu vào AI','response_content':'Kết quả AI','model_name':'Mô hình sử dụng','key':'Khóa yêu cầu','fingerprint':'Dấu vân tay nội dung'}
model_tree=ast.parse((ROOT/'backend/app/models.py').read_text(encoding='utf-8'))
schemas=[]
for cls in model_tree.body:
    if not isinstance(cls,ast.ClassDef):continue
    tab=None;fields=[]
    for node in cls.body:
        if not isinstance(node,ast.Assign):continue
        name=node.targets[0].id if isinstance(node.targets[0],ast.Name) else ''
        if name=='__tablename__':tab=ast.literal_eval(node.value)
        if isinstance(node.value,ast.Call) and isinstance(node.value.func,ast.Name) and node.value.func.id=='Column':
            call=node.value;typ=ast.unparse(call.args[0]);typ=typ.replace('String(','VARCHAR(').replace('Integer','INTEGER').replace('DateTime','DATETIME').replace('Date','DATE').replace('Boolean','BOOLEAN').replace('Text','TEXT').replace('Float','FLOAT').replace('Money()','DECIMAL(18,0)')
            kw={x.arg:ast.literal_eval(x.value) for x in call.keywords if isinstance(x.value,ast.Constant)}
            flags=[]
            if kw.get('primary_key'):flags.append('PK')
            if kw.get('unique'):flags.append('UK')
            for arg in call.args[1:]:
                if isinstance(arg,ast.Call) and getattr(arg.func,'id','')=='ForeignKey':flags.append('FK '+ast.literal_eval(arg.args[0]))
            nullable=kw.get('nullable',not kw.get('primary_key',False))
            fields.append((name,typ,'Có' if nullable else 'Không',', '.join(flags) or '—',field_desc.get(name,name)))
    if tab:schemas.append({'table':tab,'fields':fields})
for i,schema in enumerate(schemas,1):
    sub(f'4.2.{i} Bảng {schema["table"]}')
    p(descriptions[schema['table']])
    table(['Trường','Kiểu','NULL','Khóa','Ý nghĩa'],schema['fields'],[3.3,3,1.5,3.2,5],f'Cấu trúc bảng {schema["table"]}',10.5)
h('4.3 Ràng buộc và vòng đời dữ liệu',2)
p('username, mã tour, số điện thoại khách và mã booking có ràng buộc duy nhất trong mô hình; email người dùng không có UNIQUE. booking_request_keys dùng key dạng chuỗi làm khóa chính, vì vậy không thể khẳng định mọi bảng đều có id tự tăng. payments.created_by lưu chuỗi và ai_session_logs chưa liên kết khóa ngoại với users. [TL06, MK01]')
p('Kiểm tra số chỗ, tiền, thang điểm và vai trò được thực hiện tại API/CRUD; chưa có đầy đủ CHECK ở schema cho mọi miền giá trị. Bản vẽ là tài liệu mô hình tổng quan; bảng đặc tả và mã mô hình cung cấp chi tiết phục vụ triển khai. Việc sửa cấu trúc dữ liệu cần được đi kèm migration, sao lưu và kiểm thử.')
table(['Đối tượng','Trạng thái','Diễn giải'],[
('Booking','WAITING_CONFIRM','Khách gửi yêu cầu, chưa giữ chỗ.'),('Booking','PENDING','Đơn đã duyệt/đặt trực tiếp, chưa đạt mức cọc.'),('Booking','DEPOSITED','Đạt mức cọc mặc định nhưng chưa đủ tổng tiền.'),('Booking','CONFIRMED','Đã thanh toán đủ.'),('Booking','COMPLETED / CANCELLED','Hoàn thành hoặc hủy theo luồng xử lý.'),
('Departure','OPEN / FULL','Mở bán hoặc đã hết chỗ.'),('Departure','RUNNING / ON_TOUR','Các giá trị được dùng để nhận diện chuyến đang vận hành.'),('Departure','COMPLETED / CANCELLED','Chuyến kết thúc hoặc bị hủy.'),('TourGuide','AVAILABLE / ON_TOUR / OFF','Sẵn sàng, đang dẫn tour hoặc nghỉ.')],[2.8,5,8.2],'Vòng đời các đối tượng nghiệp vụ',11)

h('CHƯƠNG 5 TÍCH HỢP AI VÀ GIAO DIỆN',page=True)
h('5.1 Tư vấn tour dựa trên dữ liệu hệ thống',2)
p('Đầu vào tư vấn là nhu cầu ngôn ngữ tự nhiên của khách, ví dụ ngân sách, số ngày và sở thích. Hệ thống lấy danh sách tour/chuyến đang mở còn chỗ làm ngữ cảnh; dịch vụ AI hỗ trợ lựa chọn. Kết quả được giới hạn bởi ID hợp lệ, còn thông tin giá, ngày và số ghế hiển thị được dựng từ cơ sở dữ liệu. Cách tổ chức này giảm nguy cơ mô hình tự tạo dữ liệu giao dịch nhưng vẫn cần xử lý kết quả không hợp lệ và dữ liệu vừa thay đổi. [TL02, TL03, TL04]')
p('Danh sách gợi ý không phải cam kết đã giữ chỗ. Chỗ còn phải được kiểm tra lại ở thời điểm duyệt hoặc đặt trực tiếp. Số chỗ của nhiều ngày khởi hành không được cộng thành sức chứa của một chuyến. Khi không có khóa cloud hoặc dịch vụ không sẵn sàng, phiên bản nội bộ hỗ trợ tiếp tục các chức năng phù hợp. [TL07]')
h('5.2 Tạo mô tả và lịch trình',2)
p('Tên tour, điểm đến, thời lượng và điểm nổi bật là đầu vào tạo bản nháp. Đầu ra dùng để điền mô tả và lịch trình tóm tắt, sau đó Admin kiểm tra địa danh, thời gian, dịch vụ và tính khả thi trước khi lưu. Việc tạo được một đoạn văn không chứng minh tour đã được tạo trong cơ sở dữ liệu. [TL04, TL07]')
h('5.3 Tóm tắt phản hồi',2)
p('Chức năng phân tích thu thập nhận xét của tour, tính các chỉ số đánh giá và trình bày nội dung khen, vấn đề cần cải thiện, nhận xét về hướng dẫn viên và đề xuất hành động. Nhân viên đối chiếu với phản hồi gốc, ưu tiên xử lý vấn đề cụ thể; không coi văn bản AI là kết luận nghiệm thu chất lượng dịch vụ. Trường hợp chưa có nhận xét cần được thông báo rõ thay vì tạo nội dung đánh giá giả. [TL02, TL03, TL07]')
h('5.4 Tổ chức giao diện và thông tin phản hồi',2)
p('Thanh điều hướng phân nhóm khám phá tour, nghiệp vụ nhân viên, tài chính và quản trị. Biểu mẫu hiển thị các trường cần nhập, trạng thái và phản hồi thao tác. Mã đơn, số chỗ, tiền đã đóng và công nợ cần được giữ nhất quán giữa các màn hình; màu sắc hỗ trợ nhận biết nhưng tên trạng thái vẫn là thông tin chính. [TL06, TL07]')
with ZipFile(ROOT/'docs'/SOURCES[6]['name']) as z:
    guide_xml=E.fromstring(z.read('word/document.xml'))
    guide_rels=E.fromstring(z.read('word/_rels/document.xml.rels'))
    relationship_map={x.get('Id'):x.get('Target') for x in guide_rels}
    guide_media={rid:z.read('word/'+target) for rid,target in relationship_map.items() if target and target.startswith('media/')}
guide_body=guide_xml.find('w:body',NS)
gb=list(guide_body)
first_img=guide_xml.xpath('//a:blip',namespaces=NS)[0].get('{'+NS['r']+'}embed')
figure(guide_media[first_img],'Trang chủ VietJourney AI với dữ liệu minh họa',height=14,source='TL07')
p('Ảnh giao diện tại Chương 7 trích từ TL07, chụp bằng dữ liệu demo độc lập. Tên, ngày, tiền và số chỗ trong ảnh chỉ minh họa cách sử dụng, không phải số liệu vận hành hoặc kết quả kiểm thử hiệu năng.')

h('CHƯƠNG 6 KIỂM THỬ VÀ ĐÁNH GIÁ',page=True)
h('6.1 Môi trường và phạm vi ghi nhận',2)
p('TL05 ghi môi trường kiểm thử thủ công trên máy AMD Ryzen, RAM 8 GB, dung lượng 360 GB, kiến trúc 64 bit và Windows 11; phiên bản công cụ phát triển/SQLite chưa được ghi nhận đầy đủ. Đây là mô tả môi trường lịch sử, không phải cấu hình bắt buộc hoặc số đo trên máy hiện tại. Dữ liệu thử gồm tài khoản demo, tour, chuyến, booking, khoản thu và nhận xét. [TL05]')
p('Hồ sơ có chín nhóm kiểm thử với 49 tình huống. Phần mô tả ghi 46 tình huống Đạt và 3 Chưa đạt. Bảng kết quả ngày 25/9 ghi 45 Pass, 3 Fail và một nhãn Pas ở nhóm hàng chờ; nhãn này được giữ nguyên để đối chiếu. Chưa có bằng chứng chạy lại toàn bộ các ca cho bản web bổ sung ngày 29/10. [TL05, TL07]')
groups=['Đăng nhập và phân quyền','Quản lý tour và AI','Hướng dẫn viên','Hàng chờ tư vấn','Đặt chỗ và khách hàng','Lịch khởi hành','Kế toán và giao dịch','Thống kê','Phản hồi và AI']
summary=[]
for i,name in enumerate(groups):
    definitions=SOURCES[4]['tables'][2+i][1:]
    results=SOURCES[4]['tables'][11+i][1:]
    counts=Counter(r[3].strip() for r in results)
    summary.append((name,len(definitions),counts.get('Pass',0),counts.get('Fail',0),'; '.join(f'{k}: {v}' for k,v in counts.items() if k not in ['Pass','Fail']) or '—'))
table(['Nhóm','Số ca','Pass','Fail','Nhãn khác'],summary,[7.4,1.8,1.8,1.8,3.2],'Tổng hợp kết quả lịch sử ngày 25/9',11)
h('6.2 Các lỗi lịch sử cần đối chiếu',2)
table(['Mã trong nhóm','Nội dung ghi nhận','Cách kiểm chứng lại'],[
('Đăng nhập – TC08','Chưa chặn quyền triệt để.','Dùng token Accountant gọi API Admin; kiểm tra phản hồi 403 và dữ liệu không đổi; kiểm tra cả giao diện.'),
('Đặt chỗ – TC06','Chưa có cột phương thức thanh toán theo yêu cầu thử.','Đối chiếu cột ở màn hình đặt chỗ với sổ giao dịch; xác định vị trí thông tin theo thiết kế hiện tại trước khi nghiệm thu.'),
('Lịch khởi hành – TC02','Chưa thao tác được mở đợt khởi hành mới.','Tạo chuyến bằng Admin, kiểm tra bản ghi lưu, ngày, sức chứa và quyền; trường hợp thiếu dữ liệu phải báo lỗi.')],[3.5,5.5,7],'Ba vấn đề Fail được giữ từ hồ sơ gốc',11)
h('6.3 Kế hoạch kiểm thử cho phần cập nhật',2)
table(['Luồng','Tình huống cần xác nhận'],[
('Tra cứu đơn','Mã đúng/sai; liên hệ đúng/sai; liên kết token hợp lệ/không hợp lệ; chỉ lộ dữ liệu đúng đơn.'),
('Email và VietQR','SMTP chưa cấu hình/lỗi; APP_BASE_URL đúng; tiền và nội dung chuyển khoản phù hợp; tiền chỉ đổi khi được ghi nhận.'),
('Quản trị đặc biệt','Hủy khẩn cấp, gửi lặp, giao dịch hoàn, trạng thái HDV; xóa booking/tour có giao dịch và dữ liệu phụ thuộc.'),
('Tính nhất quán','Đặt đồng thời khi ít chỗ; hủy/thu tiền đồng thời; rollback; xử lý tạo tour thành công nhưng lịch đi lỗi.'),
('Giao diện','Biểu mẫu mới, thông báo, trạng thái trống, điều hướng theo vai trò, màn hình nhỏ và bàn phím.'),
('AI và DBMS','Cloud với cấu hình thật, ID không hợp lệ và lỗi mạng; MySQL trên môi trường triển khai; số đo hiệu năng riêng.')],[3.7,12.3],'Các nhóm kiểm thử tiếp theo được đề xuất',11.5)
p('Bảng trên là kế hoạch kiểm chứng, không phải kết quả đã chạy. Chi tiết 49 tình huống và kết quả nguồn được gộp theo từng nhóm ở Phụ lục B; mã TC được bổ sung tiền tố nhóm để tránh nhầm các TC01 trùng tên giữa các phân hệ.')

h('CHƯƠNG 7 CÀI ĐẶT VÀ HƯỚNG DẪN SỬ DỤNG',page=True)
p('Chương này biên tập từ hướng dẫn sử dụng cập nhật TL07, giữ trình tự thao tác và các ảnh minh họa. Các bước hướng đến người cài đặt và từng vai trò sử dụng; phần cấu hình chỉ dành cho người phụ trách vận hành.')
# Import all operational instructions and 16 remaining screenshots from the latest guide.
started=False;section_index=0;skip_caption=False;hardware=False
for node in gb:
    if node.tag!=qn('w:p'):continue
    text=''.join(node.xpath('.//w:t/text()',namespaces=NS)).strip()
    style=node.find('w:pPr/w:pStyle',NS)
    sid=style.get(qn('w:val')) if style is not None else ''
    if sid=='Heading1' and text=='CẤU HÌNH PHẦN CỨNG - PHẦN MỀM':
        started=True;continue
    if not started:continue
    if sid=='Heading1' and text=='CÁC CHỨC NĂNG CHÍNH':
        continue
    blips=node.xpath('.//a:blip',namespaces=NS)
    if blips:
        rid=blips[0].get('{'+NS['r']+'}embed')
        following=node.getnext()
        rawcap=''.join(following.xpath('.//w:t/text()',namespaces=NS)) if following is not None else ''
        title=re.sub(r'^Hình \d+\.\s*','',rawcap).strip() or 'Giao diện nghiệp vụ'
        figure(guide_media[rid],title,height=14,source='TL07')
        skip_caption=True;continue
    if skip_caption and text.startswith('Hình '):skip_caption=False;continue
    if not text:continue
    if sid=='Heading2':
        section_index+=1
        h(f'7.{section_index} {text}',2)
    else:
        runs=node.xpath('./w:r',namespaces=NS)
        isbold=bool(runs) and all(r.find('w:rPr/w:b',NS) is not None for r in runs if r.find('w:t',NS) is not None)
        if isbold and not text.startswith('Bước'):sub(text)
        else:p(text)

h('CHƯƠNG 8 KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN',page=True)
h('8.1 Kết quả của dự án và hồ sơ',2)
p('Hệ thống tập hợp các nghiệp vụ quản lý tour, lịch khởi hành, hướng dẫn viên, đặt chỗ, tiền và phản hồi vào cùng một ứng dụng. Hồ sơ thể hiện tám use case chính, mô hình lớp, Screen Flow và thiết kế dữ liệu; ba chức năng AI gắn với nhu cầu tư vấn và nội dung. Bản hướng dẫn cập nhật bổ sung tra cứu đơn, email, VietQR, biểu mẫu phản hồi và các thao tác quản trị.')
p('Báo cáo thống nhất bảy tài liệu theo một cấu trúc, giữ 21 sơ đồ và 17 ảnh giao diện. Yêu cầu, mô hình khái niệm, khả năng API và thao tác thực có trên giao diện được phân biệt khi cần. Phần kiểm thử giữ kết quả lịch sử và chỉ rõ những ca phải chạy lại, giúp việc đánh giá tiếp theo có căn cứ thay vì suy luận từ mô tả chức năng.')
h('8.2 Giới hạn hiện tại',2)
p('Hồ sơ chưa cung cấp số đo hiệu năng, kết quả tích hợp MySQL, xác minh đầy đủ AI cloud, SMTP và đối soát ngân hàng trên môi trường triển khai thực. Giao diện chưa có nút sửa tour trực tiếp; chức năng lọc kỳ báo cáo và xuất Excel/PDF chưa hoàn thiện. Việc nhập đánh giá chưa bắt buộc chứng minh khách đã hoàn thành chuyến. Các thao tác xóa quản trị cần chính sách lưu lịch sử và kiểm tra dữ liệu liên quan chặt chẽ hơn.')
h('8.3 Thứ tự phát triển tiếp theo',2)
steps(['Chốt lại phạm vi yêu cầu sau cập nhật: quyền quản trị, luồng hủy khẩn cấp, xóa dữ liệu, tiếp nhận phản hồi và định nghĩa số liệu báo cáo.',
       'Kiểm thử hồi quy 49 tình huống cùng các ca bổ sung về tra cứu, email, gửi lặp và tính nhất quán; cập nhật kết quả có ngày, môi trường và bằng chứng.',
       'Hoàn thiện giao diện sửa tour, điều hướng và biểu mẫu; bổ sung lọc thời gian/xuất báo cáo theo nhu cầu được xác nhận.',
       'Kiểm chứng MySQL, cloud và SMTP; thiết lập sao lưu/khôi phục, HTTPS, tài khoản và secret riêng trước khi vận hành công khai.',
       'Nghiên cứu tích hợp cổng thanh toán/đối soát tự động và kênh sử dụng riêng cho hướng dẫn viên khi có yêu cầu thực tế.'])

h('TÀI LIỆU THAM KHẢO VÀ NGUỒN NỘI BỘ',page=True)
for s in SOURCES:
    p(f'[{s["id"]}] Nhóm 17. {s["name"]}. Bản tài liệu trong thư mục docs được dùng để tổng hợp báo cáo.')
p('[MK01] Mã nguồn dự án VietJourney AI: backend/app/models.py và money.py; được đối chiếu để hoàn thiện từ điển dữ liệu. Mã nguồn được sử dụng làm bằng chứng hiện thực, không thay thế kết quả kiểm thử.')
p('Ngày khóa bản tổng hợp: 29/10/2026. Các tài liệu gốc được giữ nguyên; báo cáo được lưu thành tệp mới. Các tài liệu có mô tả phiên bản hoặc kết quả lịch sử được sử dụng trong đúng bối cảnh thời gian của chúng.')

h('PHỤ LỤC A BIỂU ĐỒ HÀNH VI CÁC USE CASE',page=True)
p('Phụ lục giữ nguyên các ảnh Activity và Sequence đã được nhóm cung cấp và chấp nhận sử dụng. Mỗi cặp minh họa một use case; mô tả và giới hạn phiên bản ở Chương 2 được dùng kèm để đọc sơ đồ. [TL03]')
for i,(code,title,*_) in enumerate(ucs):
    h(f'A.{i+1} {code} {title}',2,page=i>0,toc=False)
    figure(DIAGRAMS[6+2*i],f'{code} Activity',height=16.6,source='TL03')
    figure(DIAGRAMS[7+2*i],f'{code} Sequence',height=5.1,source='TL03')

h('PHỤ LỤC B TÌNH HUỐNG VÀ KẾT QUẢ KIỂM THỬ',page=True)
p('Mỗi dòng kết hợp một tình huống ở phần mô tả và kết quả lịch sử tương ứng tại TL05. Dữ liệu, điều kiện và mong đợi được giữ để đối chiếu; lỗi chính tả trình bày không làm thay đổi trạng thái. Tiền tố G01–G09 phân biệt các nhóm có cùng mã TC. Các quyền và số liệu mẫu là nội dung của ca nguồn, cần cập nhật khi chạy lại phiên bản mới.')
case_records=[]
for i,name in enumerate(groups):
    h(f'B.{i+1} {name}',2)
    defs=SOURCES[4]['tables'][2+i][1:]
    results={r[0]:r for r in SOURCES[4]['tables'][11+i][1:]}
    rows=[]
    for rr in defs:
        result=results[rr[0]]
        cid=f'G{i+1:02d}-{rr[0]}'
        scenario=f'{rr[1]}\n{rr[2]}\nĐiều kiện: {rr[3]}\nDữ liệu/thao tác: {rr[4]}'
        status=f'{result[3]}\nNgày: {result[1]}\n{result[2]}'
        rows.append((cid,scenario,rr[5],status))
        case_records.append({'id':cid,'source_id':rr[0],'group':name,'expected':rr[5],'status':result[3],'date':result[1],'tester':result[2]})
    table(['Mã','Tình huống và đầu vào','Kết quả mong đợi','Ghi nhận gốc'],rows,[2.1,6.4,5.2,2.3],f'Ca kiểm thử nhóm {i+1} {name}',10.5)
note('Nhóm hàng chờ có một kết quả ghi Pas trong nguồn; giữ nguyên nhãn và cần xác nhận khi chuẩn hóa hồ sơ. Ba kết quả Fail không được đổi thành Pass chỉ dựa vào ảnh hoặc tài liệu mô tả.')

# TOC uses live page references; fill cached values after the first render.
toc_entries=[x for x in headings if x['toc'] and x['level']<=2]
for toc_index,item in enumerate(toc_entries):
    if toc_index==28:
        continuation=d.add_paragraph('MỤC LỤC (tiếp theo)')
        continuation.alignment=WD_ALIGN_PARAGRAPH.CENTER
        continuation.paragraph_format.page_break_before=True
        continuation.runs[0].bold=True
        toc_marker._p.addprevious(continuation._p)
    para=d.add_paragraph()
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.0
    if item['level']==2:para.paragraph_format.left_indent=Cm(0.5)
    para.paragraph_format.tab_stops.add_tab_stop(Cm(15.9),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('w:anchor'),item['anchor'])
    rr=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');rp.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'23');rp.append(sz)
    if item['level']==1:rp.append(OxmlElement('w:b'))
    rr.append(rp);tt=OxmlElement('w:t');tt.text=item['text'];rr.append(tt);rr.append(OxmlElement('w:tab'));hyperlink.append(rr)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),f' PAGEREF {item["anchor"]} \\h ')
    field.set(qn('w:dirty'),'true');run=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text='–';run.append(tx);field.append(run);hyperlink.append(field)
    para._p.append(hyperlink)
    toc_marker._p.addprevious(para._p)
toc_marker._element.getparent().remove(toc_marker._element)

d.save(OUT)
with ZipFile(OUT) as z:assert z.testzip() is None
assert len(figures)==38,len(figures)
assert len(case_records)==49
assert len(schemas)==10
for s in SOURCES:assert hashlib.sha256((ROOT/'docs'/s['name']).read_bytes()).hexdigest()==s['sha256']
manifest={'output':str(OUT),'sources':[{k:s[k] for k in ['id','name','sha256']} for s in SOURCES],'headings':headings,'figures':figures,'tables':tables,'cases':case_records,'schemas':schemas}
(QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'artifact.md').write_text('''# Combined report specification
Audience: lecturer and project team; one academic Vietnamese report synthesized from seven DOCX files.
Typography: Times New Roman 13 body, 16/14 headings, black. A4 portrait, left3/right2/top2/bottom2 cm. Same family and page system as supplied project documents.
Structure: cover; linked TOC; abstract; eight chapters; internal references; appendix of sixteen Activity/Sequence images; full forty-nine historical test cases.
Figures: preserve all21 accepted diagrams and17 screenshots from the latest user guide. Original image bytes retained.
Evidence: source hashes retained. Requirements/implementation/history distinguished; no invented university, lecturer, measured performance, external survey data or new test outcomes.
Tables: fixed widths16cm; repeated headers; no row splitting; 10.5–12pt text; light borders, neutral shading.
Render: native extracted LibreOffice and bundled Poppler; inspect all final page PNGs and verify content bounds and cached TOC.
''',encoding='utf-8')
print(json.dumps({'output':str(OUT),'figures':len(figures),'tables':len(tables),'test_cases':len(case_records),'toc_entries':len(toc_entries)},ensure_ascii=False))
