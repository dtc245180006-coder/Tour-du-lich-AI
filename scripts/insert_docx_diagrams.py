"""Insert accepted diagrams with local OOXML edits; preserve source formatting."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from io import BytesIO
import hashlib
import json
from lxml import etree as E
from PIL import Image
from docx import Document
from docx.shared import Emu

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/DOCX_DA_CHINH_SUA'
OUT = ROOT / 'docs/DOCX_DA_CHEN_SO_DO'
OUT.mkdir(exist_ok=True)
NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
}
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
images = {int(p.name.split('.')[0]): p for p in (ROOT/'docs/diagrams').glob('*.jpg')}
assert set(images) == set(range(1, 22))
manifest = []

def xml_bytes(root):
    return E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

for source in sorted(SOURCE.glob('*.docx')):
    with ZipFile(source) as archive:
        parts = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    root = E.fromstring(parts['word/document.xml'])
    rels = E.fromstring(parts['word/_rels/document.xml.rels'])
    paragraphs = root.xpath('//w:p', namespaces=NS)
    drawings = root.xpath('//w:drawing', namespaces=NS)
    numbers = []
    if source.name.startswith('04_'):
        numbers = [3]
    elif source.name.startswith('06_'):
        numbers = [5, 4]
    elif source.name.startswith('Gen3'):
        numbers = [1, 7, 6] + list(range(8, 22))
    assert len(drawings) == len(numbers), (source.name, len(drawings), numbers)
    section = root.xpath('//w:sectPr', namespaces=NS)[-1]
    size = section.find('w:pgSz', NS)
    margins = section.find('w:pgMar', NS)
    width = (int(size.get('{%s}w'%NS['w'])) - int(margins.get('{%s}left'%NS['w'])) - int(margins.get('{%s}right'%NS['w']))) * 635
    height = (int(size.get('{%s}h'%NS['w'])) - int(margins.get('{%s}top'%NS['w'])) - int(margins.get('{%s}bottom'%NS['w']))) * 635 - 457200
    if source.name.startswith('02_'):
        candidates = [p for p in paragraphs if 'Sơ đồ phân cấp chức năng' in ''.join(p.xpath('.//w:t/text()', namespaces=NS))]
        assert len(candidates) == 1
        heading = candidates[0]
        following = heading.getnext()
        if following is not None and following.tag == '{%s}p'%NS['w'] and not following.xpath('.//w:t|.//w:drawing', namespaces=NS):
            target = following
        else:
            target = E.Element('{%s}p'%NS['w'])
            heading.addnext(target)
        temp = Document()
        temp.add_picture(str(images[2]), width=Emu(width))
        drawing = E.fromstring(E.tostring(temp.element.xpath('//w:drawing')[0]))
        run = E.SubElement(target, '{%s}r'%NS['w'])
        run.append(drawing)
        drawings = [drawing]
        numbers = [2]
        rid = 'rIdAcceptedDiagram2'
        drawing.xpath('.//a:blip', namespaces=NS)[0].set('{%s}embed'%NS['r'], rid)
        E.SubElement(rels, '{%s}Relationship'%REL, Id=rid, Type=NS['r']+'/image', Target='media/diagram_02.jpg')
    for drawing, number in zip(drawings, numbers):
        picture = images[number]
        blob = picture.read_bytes()
        with Image.open(BytesIO(blob)) as im:
            iw, ih = im.size
        extent = drawing.xpath('.//wp:extent', namespaces=NS)[0]
        old_width = int(extent.get('cx'))
        scale = min(min(old_width, width)/iw, height/ih)
        cx, cy = round(iw*scale), round(ih*scale)
        extent.set('cx', str(cx)); extent.set('cy', str(cy))
        for node in drawing.xpath('.//a:xfrm/a:ext', namespaces=NS):
            node.set('cx', str(cx)); node.set('cy', str(cy))
        for crop in drawing.xpath('.//a:srcRect', namespaces=NS):
            crop.getparent().remove(crop)
        props = drawing.xpath('.//wp:docPr', namespaces=NS)[0]
        props.set('name', picture.stem)
        props.set('descr', picture.stem)
        if source.name.startswith('02_'):
            props.set('id', str(max([int(x.get('id')) for x in root.xpath('//wp:docPr', namespaces=NS)] + [0])+1))
        rid = drawing.xpath('.//a:blip', namespaces=NS)[0].get('{%s}embed'%NS['r'])
        relation = rels.xpath('*[@Id=$rid]', rid=rid)[0]
        relation.set('Target', f'media/diagram_{number:02d}.jpg')
        parts[f'word/media/diagram_{number:02d}.jpg'] = blob
        manifest.append({'document': source.name, 'image': picture.name, 'cx': cx, 'cy': cy, 'sha256': hashlib.sha256(blob).hexdigest()})
    if numbers:
        parts['word/document.xml'] = xml_bytes(root)
        parts['word/_rels/document.xml.rels'] = xml_bytes(rels)
        types = E.fromstring(parts['[Content_Types].xml'])
        if not types.xpath('*[@Extension="jpg"]'):
            E.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Default', Extension='jpg', ContentType='image/jpeg')
        parts['[Content_Types].xml'] = xml_bytes(types)
    with ZipFile(OUT/source.name, 'w', ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    # Verify source text, font/paragraph properties, tables, and sections remain intact.
    with ZipFile(source) as before, ZipFile(OUT/source.name) as after:
        a = E.fromstring(before.read('word/document.xml'))
        b = E.fromstring(after.read('word/document.xml'))
        for xpath in ['//w:t', '//w:rPr', '//w:pPr', '//w:tblPr', '//w:sectPr']:
            assert [E.tostring(x) for x in a.xpath(xpath, namespaces=NS)] == [E.tostring(x) for x in b.xpath(xpath, namespaces=NS)], (source.name, xpath)
        for name in before.namelist():
            if name not in ['word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml']:
                assert before.read(name) == after.read(name), (source.name, name)
    Document(OUT/source.name)  # Package and relationships are readable.
assert len(manifest) == 21
(OUT/'DOI_CHIEU_ANH.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
with ZipFile(ROOT/'docs/BO_DOCX_DA_CHEN_SO_DO.zip', 'w', ZIP_DEFLATED) as archive:
    for path in OUT.glob('*.docx'):
        archive.write(path, path.name)
print(json.dumps({'documents': 6, 'inserted_images': len(manifest), 'output': str(OUT)}, ensure_ascii=False))
